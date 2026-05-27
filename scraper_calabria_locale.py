#!/usr/bin/env python3
"""
Monitoraggio Atti Calabria - Generatore Excel
Portali: ASP CO/CZ/KR/RC/VV, Regione Calabria, TAR Catanzaro

Questo script può essere eseguito sia in cloud (solo Regione Calabria funziona)
sia in locale su Windows dove tutti i portali sono raggiungibili.

Output: Monitoraggio_Atti_Calabria_DD-MM-YYYY.xlsx
"""

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re
import time
import datetime
import os
import sys
import traceback
import urllib3
from urllib.parse import urljoin, urlencode

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ─────────────────────────────────────────────────────────────────────────────
# DATE
# ─────────────────────────────────────────────────────────────────────────────
TODAY            = datetime.date.today()
DATE_FROM        = TODAY - datetime.timedelta(days=2)
DATE_FROM_TAR    = TODAY - datetime.timedelta(days=3)

DATE_FROM_STR    = DATE_FROM.strftime("%Y-%m-%d")
DATE_TODAY_STR   = TODAY.strftime("%Y-%m-%d")
DATE_FROM_TAR_STR= DATE_FROM_TAR.strftime("%Y-%m-%d")

OUTPUT_FNAME     = f"Monitoraggio_Atti_Calabria_{TODAY.strftime('%d-%m-%Y')}.xlsx"
# Percorso di output - modifica se necessario
_KDIR = r'C:\Users\aoedo\kDrive\01_Lavoro\Clienti\Starbene\Atti'
OUTPUT_PATH = os.path.join(_KDIR if os.path.isdir(_KDIR) else os.path.dirname(os.path.abspath(__file__)), OUTPUT_FNAME)

# ─────────────────────────────────────────────────────────────────────────────
# KEYWORDS
# ─────────────────────────────────────────────────────────────────────────────
KEYWORDS = [
    # "ADI" ha word boundary dedicato sotto (evita match in "CALABRIADIP" ecc.)
    "Assistenza domiciliare",
    "ANMIC",
    "Accreditamento",
    "Aumento di budget",
    "Autismo",
    "Autorizzazione all'esercizio",
    "Autorizzazione alla realizzazione",
    "Autorizzazioni",
    "Budget",
    "Casa Giardino",
    "Centro San Giuseppe",
    "Centro salute e benessere",
    "Fabbisogni LEA",
    "Fisiolab",
    "Fisioterapia",
    "Parere commissione",
    "Presa d'atto verifica",
    "Programmazione",
    "Rete riabilitativa",
    "Rete territoriale",
    "Riabilitazione estensiva",
    "Riconversione prestazioni",
    "Rinnovo accreditamento",
    "San Teodoro",
    "Savelli Hospital",
    "Starbene",
    "Verifica requisiti",
    "Villa San Giuseppe",
    "Villa del Rosario",
]
ADI_RE         = re.compile(r'\bADI\b',  re.IGNORECASE)  # word boundary per evitare false positives
LIFE_RE        = re.compile(r'\bLIFE\b', re.IGNORECASE)
KEYWORDS_RE    = [re.compile(re.escape(kw), re.IGNORECASE) for kw in KEYWORDS]

def matches_keywords(text: str) -> bool:
    if not text:
        return False
    if ADI_RE.search(text):
        return True
    if LIFE_RE.search(text):
        return True
    for pattern in KEYWORDS_RE:
        if pattern.search(text):
            return True
    return False

# ─────────────────────────────────────────────────────────────────────────────
# SESSION HTTP
# ─────────────────────────────────────────────────────────────────────────────
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept":          "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "it-IT,it;q=0.9,en-US;q=0.8,en;q=0.7",
    "Accept-Encoding": "gzip, deflate",
    "Connection":      "keep-alive",
    "Cache-Control":   "max-age=0",
}

def make_session() -> requests.Session:
    s = requests.Session()
    s.headers.update(HEADERS)
    s.verify = False
    return s


# ─────────────────────────────────────────────────────────────────────────────
# 1. SCRAPER ASP SISR (Albo Online)
# ─────────────────────────────────────────────────────────────────────────────

ASP_PORTALS = [
    ("ASP Cosenza",         "https://online-aspco.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
    ("ASP Catanzaro",       "https://online-aspcz.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
    ("ASP Crotone",         "https://online-aspkr.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
    ("ASP Reggio Calabria", "https://online-asprc.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
    ("ASP Vibo Valentia",   "https://online-aspvv.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
]

def scrape_asp(name: str, base_url: str) -> tuple[list[dict], str | None]:
    """
    Scrape ASP Albo Online.
    Ritorna (lista_match, errore_o_None)
    """
    results = []
    error   = None
    session = make_session()
    print(f"\n{'='*60}\n  Scraping {name}\n  URL: {base_url}")

    # ── Sessione iniziale ──────────────────────────────────────────
    try:
        resp = session.get(base_url, timeout=30, verify=False)
        resp.raise_for_status()
    except requests.exceptions.Timeout:
        error = f"Timeout: il portale non ha risposto entro 30s"
        print(f"  [ERRORE] {error}")
        return results, error
    except requests.exceptions.HTTPError as e:
        error = f"HTTP {e.response.status_code}: {e.response.reason}"
        print(f"  [ERRORE] {error}")
        return results, error
    except Exception as e:
        error = f"{type(e).__name__}: {e}"
        print(f"  [ERRORE] {error}")
        return results, error

    soup = BeautifulSoup(resp.text, "lxml")

    # Raccoglie hidden fields (CSRF, viewstate, ecc.)
    hidden_fields = {}
    for inp in soup.find_all("input", type="hidden"):
        nm = inp.get("name"); vl = inp.get("value", "")
        if nm:
            hidden_fields[nm] = vl

    # Parametri di ricerca
    search_params = {
        **hidden_fields,
        "dataPubblicazioneDal": DATE_FROM_STR,
        "dataPubblicazioneAl":  DATE_TODAY_STR,
        "oggetto": "",
        "numero":  "",
    }

    # Form action e method
    form        = soup.find("form")
    action_url  = base_url
    form_method = "get"
    if form:
        action = form.get("action", "")
        if action:
            action_url = urljoin(base_url, action)
        form_method = form.get("method", "get").lower()
        for inp in form.find_all(["input", "select"]):
            nm = inp.get("name"); vl = inp.get("value", ""); tp = inp.get("type", "text")
            if nm and nm not in hidden_fields and tp not in ("submit", "button", "image", "reset"):
                search_params[nm] = vl

    print(f"  Form action={action_url}, method={form_method}")
    print(f"  Ricerca: {DATE_FROM_STR} → {DATE_TODAY_STR}")

    # ── Paginazione ────────────────────────────────────────────────
    page = 1
    total = 0

    while True:
        print(f"  Pagina {page}...", end=" ", flush=True)
        params = dict(search_params)
        if page > 1:
            for k in ("page", "pagina", "currentPage", "paginaCorrente", "pageNumber"):
                params[k] = page

        try:
            if form_method == "post":
                r = session.post(action_url, data=params, timeout=30, verify=False)
            else:
                r = session.get(action_url, params=params, timeout=30, verify=False)
            r.raise_for_status()
        except Exception as e:
            print(f"[ERRORE pagina {page}] {e}")
            break

        page_soup = BeautifulSoup(r.text, "lxml")
        table = (
            page_soup.find("table", id=re.compile(r"result|albo|atti", re.I))
            or page_soup.find("table", class_=re.compile(r"result|albo|atti|table", re.I))
            or page_soup.find("table")
        )
        if not table:
            print("[nessuna tabella trovata]")
            break

        rows = [tr for tr in table.find_all("tr") if tr.find("td")]
        th_rows = [tr for tr in table.find_all("tr") if tr.find("th") and not tr.find("td")]
        headers = []
        if th_rows:
            headers = [c.get_text(strip=True) for c in th_rows[0].find_all("th")]

        print(f"{len(rows)} righe", end=" ", flush=True)
        if not rows:
            break

        page_match = 0
        for row in rows:
            cells = row.find_all(["th", "td"])
            cell_texts = [c.get_text(strip=True) for c in cells]
            link_url = ""
            for cell in cells:
                a = cell.find("a", href=True)
                if a:
                    link_url = urljoin(base_url, a["href"])
                    break
            row_dict = {(headers[j] if j < len(headers) else f"col{j}"): t
                        for j, t in enumerate(cell_texts)}
            oggetto  = (row_dict.get("Oggetto") or row_dict.get("Descrizione")
                        or row_dict.get("Titolo") or " ".join(cell_texts[2:4]))
            data_row = (row_dict.get("Data pubblicazione") or row_dict.get("Data Pubblicazione")
                        or row_dict.get("Data") or (cell_texts[0] if cell_texts else ""))
            if matches_keywords(oggetto):
                results.append({"data": data_row, "oggetto": oggetto, "link": link_url, "raw": row_dict})
                page_match += 1

        total += len(rows)
        print(f"→ {page_match} match")

        # Cerca pagina successiva
        next_pg = _find_asp_next_page(page_soup, page)
        if next_pg is None:
            break
        page = next_pg
        time.sleep(0.4)

    print(f"  Tot. righe: {total}, match: {len(results)}")
    return results, error


def _find_asp_next_page(soup: BeautifulSoup, current: int) -> int | None:
    for a in soup.find_all("a"):
        txt = a.get_text(strip=True)
        if txt in (">", "»", "›", "Successiva", "Next", "Avanti"):
            href = a.get("href", "")
            if href and href != "#":
                m = re.search(r"[pP]age[=_]?(\d+)|[pP]agina[=_]?(\d+)|currentPage[=]?(\d+)", href)
                if m:
                    return int(next(x for x in m.groups() if x))
                return current + 1
    pg = soup.find(class_=re.compile(r"paginat|pager|pagination", re.I))
    if pg:
        for a in pg.find_all("a"):
            if a.get_text(strip=True).isdigit() and int(a.get_text(strip=True)) == current + 1:
                return current + 1
    return None


# ─────────────────────────────────────────────────────────────────────────────
# 2. SCRAPER REGIONE CALABRIA
# ─────────────────────────────────────────────────────────────────────────────

REGIONE_BASE = "https://www.regione.calabria.it/provvedimenti-della-regione/"

def scrape_regione() -> tuple[list[dict], str | None]:
    """
    Struttura tabella: Tipologia | Data Repertoriazione | N (th) | Dipartimento | Oggetto | →(link)
    Paginazione: ?paged=N&filter_date_from=YYYY-MM-DD&filter_date_to=YYYY-MM-DD
    """
    results = []
    error   = None
    session = make_session()
    session.headers["Referer"] = "https://www.regione.calabria.it/"
    print(f"\n{'='*60}\n  Scraping Regione Calabria\n  URL: {REGIONE_BASE}")
    print(f"  Ricerca: {DATE_FROM_STR} → {DATE_TODAY_STR}")

    page       = 1
    total      = 0
    base_params= {"filter_date_from": DATE_FROM_STR, "filter_date_to": DATE_TODAY_STR}

    while True:
        print(f"  Pagina {page}...", end=" ", flush=True)
        params = dict(base_params)
        if page > 1:
            params["paged"] = page

        try:
            r = session.get(REGIONE_BASE, params=params, timeout=30, verify=False)
            r.raise_for_status()
        except Exception as e:
            err_msg = f"{type(e).__name__}: {e}"
            print(f"[ERRORE] {err_msg}")
            error = err_msg
            break

        soup = BeautifulSoup(r.text, "lxml")
        table = soup.find("table")
        if not table:
            print("[nessuna tabella]")
            break

        rows = [tr for tr in table.find_all("tr") if tr.find("td")]
        print(f"{len(rows)} righe", end=" ", flush=True)
        if not rows:
            break

        page_match = 0
        for row in rows:
            cells = row.find_all(["th", "td"])
            if len(cells) < 4:
                continue
            # Struttura: [0]=Tipologia [1]=Data [2]=N(th) [3]=Dipartimento [4]=Oggetto [5]=Link
            data_row    = cells[1].get_text(strip=True)
            dipartimento= cells[3].get_text(strip=True) if len(cells) > 3 else ""
            oggetto     = cells[4].get_text(strip=True) if len(cells) > 4 else cells[2].get_text(strip=True)
            link_url    = ""
            if len(cells) > 5:
                a = cells[5].find("a", href=True)
                if a:
                    link_url = a["href"]
                    # link già assoluto di solito
            else:
                # cerca link in tutta la riga
                for c in cells:
                    a = c.find("a", href=True)
                    if a and a["href"] != "#":
                        link_url = a["href"]
                        break

            # Filtra SOLO su Oggetto (spec: "filtra la colonna Oggetto")
            if matches_keywords(oggetto):
                results.append({
                    "data":         data_row,
                    "oggetto":      oggetto,
                    "dipartimento": dipartimento,
                    "link":         link_url,
                })
                page_match += 1

        total += len(rows)
        print(f"→ {page_match} match")

        # Cerca next page nel paginatore WordPress
        pag = soup.find("ul", class_="page-numbers")
        has_next = False
        if pag:
            for a in pag.find_all("a"):
                if "next" in a.get("class", []):
                    has_next = True
                    break
        if not has_next:
            break
        page += 1
        time.sleep(0.5)

    print(f"  Tot. righe: {total}, match: {len(results)}")
    return results, error


# ─────────────────────────────────────────────────────────────────────────────
# 3. SCRAPER TAR CATANZARO
# ─────────────────────────────────────────────────────────────────────────────

TAR_BASE = "https://www.giustizia-amministrativa.it/provvedimenti-tar-catanzaro"

def scrape_tar() -> tuple[list[dict], str | None]:
    results = []
    error   = None
    session = make_session()
    session.headers["Referer"] = "https://www.giustizia-amministrativa.it/"
    print(f"\n{'='*60}\n  Scraping TAR Catanzaro\n  URL: {TAR_BASE}")
    print(f"  Ricerca: {DATE_FROM_TAR_STR} → {DATE_TODAY_STR}")

    page  = 0
    total = 0

    while True:
        print(f"  Pagina {page+1}...", end=" ", flush=True)
        params = {
            "publishDateFrom": DATE_FROM_TAR_STR,
            "publishDateTo":   DATE_TODAY_STR,
            "page":  page,
            "size":  50,
        }
        try:
            r = session.get(TAR_BASE, params=params, timeout=30, verify=False)
            r.raise_for_status()
        except requests.exceptions.HTTPError as e:
            err_msg = f"HTTP {e.response.status_code}: {e.response.reason}"
            print(f"[ERRORE] {err_msg}")
            error = err_msg
            break
        except Exception as e:
            err_msg = f"{type(e).__name__}: {e}"
            print(f"[ERRORE] {err_msg}")
            error = err_msg
            break

        soup  = BeautifulSoup(r.text, "lxml")
        table = soup.find("table")
        if not table:
            print("[nessuna tabella]")
            break

        rows = [tr for tr in table.find_all("tr") if tr.find("td")]
        print(f"{len(rows)} righe", end=" ", flush=True)
        if not rows:
            break

        th_header = [tr for tr in table.find_all("tr") if tr.find("th")]
        headers   = [c.get_text(strip=True) for c in th_header[0].find_all("th")] if th_header else []

        page_match = 0
        for row in rows:
            cells = row.find_all(["th", "td"])
            ctext = [c.get_text(strip=True) for c in cells]
            link_url = ""
            for c in cells:
                a = c.find("a", href=True)
                if a and a["href"] != "#":
                    link_url = urljoin(TAR_BASE, a["href"])
                    break
            rd = {(headers[j] if j < len(headers) else f"col{j}"): t
                  for j, t in enumerate(ctext)}
            parte   = (rd.get("Parte") or rd.get("parte") or rd.get("Ricorrente") or ctext[1] if len(ctext)>1 else "")
            oggetto = (rd.get("Oggetto") or rd.get("Materia") or rd.get("oggetto") or "")
            data_row= (rd.get("Data") or rd.get("DataDeposito") or ctext[0] if ctext else "")
            combined = f"{parte} {oggetto}"
            if matches_keywords(combined):
                results.append({"data": data_row, "parte": parte, "oggetto": oggetto,
                                 "link": link_url, "raw": rd})
                page_match += 1

        total += len(rows)
        print(f"→ {page_match} match")

        if not _tar_has_next(soup):
            break
        page += 1
        time.sleep(0.5)

    print(f"  Tot. righe: {total}, match: {len(results)}")
    return results, error


def _tar_has_next(soup: BeautifulSoup) -> bool:
    for a in soup.find_all("a"):
        txt  = a.get_text(strip=True)
        href = a.get("href", "")
        if txt in (">", "»", "›", "Pagina successiva", "Next", "Avanti") and href and href != "#":
            return True
    pag = soup.find(class_=re.compile(r"paginat|pager|pagination", re.I))
    if pag:
        for a in pag.find_all("a"):
            if a.get_text(strip=True) in (">", "»", "›"):
                return True
    return False


# ─────────────────────────────────────────────────────────────────────────────
# 4. EXCEL WRITER
# ─────────────────────────────────────────────────────────────────────────────

STYLE = {
    "header_fill":  PatternFill("solid", fgColor="1F4E79"),
    "header_font":  Font(bold=True, color="FFFFFF", size=11, name="Calibri"),
    "even_fill":    PatternFill("solid", fgColor="DCE6F1"),
    "link_font":    Font(color="0563C1", underline="single", name="Calibri"),
    "body_font":    Font(name="Calibri", size=10),
    "warn_fill":    PatternFill("solid", fgColor="FFF2CC"),
    "warn_font":    Font(bold=True, color="7F6000", name="Calibri"),
    "thin_border":  Border(
        left=Side(style='thin', color='BDD7EE'),
        right=Side(style='thin', color='BDD7EE'),
        top=Side(style='thin', color='BDD7EE'),
        bottom=Side(style='thin', color='BDD7EE'),
    ),
}


def _write_error_sheet(ws, portal_name: str, error_msg: str, portal_url: str):
    """Scrive un foglio con messaggio di errore quando il portale non è raggiungibile."""
    ws.column_dimensions["A"].width = 80
    ws.row_dimensions[1].height = 25

    cell = ws.cell(row=1, column=1,
                   value=f"⚠ Portale non raggiungibile dall'ambiente di esecuzione remoto")
    cell.font  = STYLE["warn_font"]
    cell.fill  = STYLE["warn_fill"]
    cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.cell(row=2, column=1, value=f"Portale: {portal_name}")
    ws.cell(row=3, column=1, value=f"URL: {portal_url}")
    ws.cell(row=3, column=1).hyperlink = portal_url
    ws.cell(row=3, column=1).font = STYLE["link_font"]
    ws.cell(row=4, column=1, value=f"Errore: {error_msg}")
    ws.cell(row=5, column=1,
            value="ℹ  Per ottenere i dati, eseguire il file 'scraper_calabria_locale.py' "
                  "direttamente sul proprio PC (Python 3.11+ con requests, bs4, openpyxl).")
    ws.cell(row=5, column=1).font = Font(italic=True, color="595959", name="Calibri")
    ws.row_dimensions[5].height = 30


def _setup_header(ws, headers: list[str]):
    ws.row_dimensions[1].height = 28
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = STYLE["header_font"]
        cell.fill      = STYLE["header_fill"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = STYLE["thin_border"]
    ws.freeze_panes = "A2"


def write_regione_sheet(ws, rows: list[dict]):
    """Foglio Regione Calabria: Data | Oggetto | Descrizione breve | Link"""
    headers = ["Data", "Oggetto", "Descrizione breve (150 car.)", "Link"]
    _setup_header(ws, headers)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 65
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 70

    if not rows:
        c = ws.cell(row=2, column=1, value="Nessun risultato nel periodo selezionato")
        c.font = Font(italic=True, color="595959")
        return

    for r_idx, row in enumerate(rows, 2):
        fill = STYLE["even_fill"] if r_idx % 2 == 0 else PatternFill()
        data_val = row.get("data",    "")
        oggetto  = row.get("oggetto", "") or ""
        descr    = (oggetto[:150] + "…") if len(oggetto) > 150 else oggetto
        link_url = row.get("link",    "") or ""

        vals = [data_val, oggetto, descr, link_url]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font      = STYLE["body_font"]
            cell.border    = STYLE["thin_border"]
            cell.alignment = Alignment(vertical="top", wrap_text=(col == 2))
            if fill.fill_type:
                cell.fill = fill
            if col == 4 and link_url:
                cell.hyperlink = link_url
                cell.font      = STYLE["link_font"]

        ws.row_dimensions[r_idx].height = 45 if len(oggetto) > 100 else 30


def write_asp_sheet(ws, rows: list[dict]):
    """Foglio ASP: Data | Oggetto | Descrizione breve | Link"""
    headers = ["Data", "Oggetto", "Descrizione breve (150 car.)", "Link"]
    _setup_header(ws, headers)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 65
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 70

    if not rows:
        c = ws.cell(row=2, column=1, value="Nessun risultato nel periodo selezionato")
        c.font = Font(italic=True, color="595959")
        return

    for r_idx, row in enumerate(rows, 2):
        fill = STYLE["even_fill"] if r_idx % 2 == 0 else PatternFill()
        data_val = row.get("data",    "")
        oggetto  = row.get("oggetto", "") or ""
        descr    = (oggetto[:150] + "…") if len(oggetto) > 150 else oggetto
        link_url = row.get("link",    "") or ""

        vals = [data_val, oggetto, descr, link_url]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font      = STYLE["body_font"]
            cell.border    = STYLE["thin_border"]
            cell.alignment = Alignment(vertical="top", wrap_text=(col == 2))
            if fill.fill_type:
                cell.fill = fill
            if col == 4 and link_url:
                cell.hyperlink = link_url
                cell.font      = STYLE["link_font"]

        ws.row_dimensions[r_idx].height = 45 if len(oggetto) > 100 else 30


def write_tar_sheet(ws, rows: list[dict]):
    """Foglio TAR: Data | Parte | Oggetto | Descrizione breve | Link"""
    headers = ["Data", "Parte", "Oggetto", "Descrizione breve (150 car.)", "Link"]
    _setup_header(ws, headers)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 70

    if not rows:
        c = ws.cell(row=2, column=1, value="Nessun risultato nel periodo selezionato")
        c.font = Font(italic=True, color="595959")
        return

    for r_idx, row in enumerate(rows, 2):
        fill = STYLE["even_fill"] if r_idx % 2 == 0 else PatternFill()
        data_val = row.get("data",    "")
        parte    = row.get("parte",   "") or ""
        oggetto  = row.get("oggetto", "") or ""
        descr    = (oggetto[:150] + "…") if len(oggetto) > 150 else oggetto
        link_url = row.get("link",    "") or ""

        vals = [data_val, parte, oggetto, descr, link_url]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font      = STYLE["body_font"]
            cell.border    = STYLE["thin_border"]
            cell.alignment = Alignment(vertical="top", wrap_text=(col == 3))
            if fill.fill_type:
                cell.fill = fill
            if col == 5 and link_url:
                cell.hyperlink = link_url
                cell.font      = STYLE["link_font"]

        ws.row_dimensions[r_idx].height = 45 if len(oggetto) > 80 else 30


# ─────────────────────────────────────────────────────────────────────────────
# 5. MAIN
# ─────────────────────────────────────────────────────────────────────────────

def main():
    print(f"\n{'#'*62}")
    print(f"  MONITORAGGIO ATTI CALABRIA")
    print(f"  Esecuzione: {TODAY.strftime('%d/%m/%Y')}")
    print(f"  Periodo ASP/Regione: {DATE_FROM_STR} → {DATE_TODAY_STR} (2 giorni)")
    print(f"  Periodo TAR:         {DATE_FROM_TAR_STR} → {DATE_TODAY_STR} (3 giorni)")
    print(f"  Output: {OUTPUT_PATH}")
    print(f"{'#'*62}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── ASP portali ──────────────────────────────────────────────────────────
    for asp_name, asp_url in ASP_PORTALS:
        ws = wb.create_sheet(title=asp_name[:31])
        try:
            data, err = scrape_asp(asp_name, asp_url)
        except Exception as e:
            data, err = [], str(e)
            traceback.print_exc()

        if err and not data:
            _write_error_sheet(ws, asp_name, err, asp_url)
        else:
            write_asp_sheet(ws, data)
        print(f"  ✓ Foglio '{asp_name}': {len(data)} match" + (f" [ERR: {err}]" if err else ""))

    # ── Regione Calabria ─────────────────────────────────────────────────────
    ws_reg = wb.create_sheet(title="Regione Calabria")
    try:
        reg_data, reg_err = scrape_regione()
    except Exception as e:
        reg_data, reg_err = [], str(e)
        traceback.print_exc()

    if reg_err and not reg_data:
        _write_error_sheet(ws_reg, "Regione Calabria", reg_err, REGIONE_BASE)
    else:
        write_regione_sheet(ws_reg, reg_data)
    print(f"  ✓ Foglio 'Regione Calabria': {len(reg_data)} match" + (f" [ERR: {reg_err}]" if reg_err else ""))

    # ── TAR Catanzaro ─────────────────────────────────────────────────────────
    ws_tar = wb.create_sheet(title="TAR Catanzaro")
    try:
        tar_data, tar_err = scrape_tar()
    except Exception as e:
        tar_data, tar_err = [], str(e)
        traceback.print_exc()

    if tar_err and not tar_data:
        _write_error_sheet(ws_tar, "TAR Catanzaro", tar_err, TAR_BASE)
    else:
        write_tar_sheet(ws_tar, tar_data)
    print(f"  ✓ Foglio 'TAR Catanzaro': {len(tar_data)} match" + (f" [ERR: {tar_err}]" if tar_err else ""))

    # ── Salva ────────────────────────────────────────────────────────────────
    os.makedirs(os.path.dirname(OUTPUT_PATH) or ".", exist_ok=True)
    wb.save(OUTPUT_PATH)

    print(f"\n{'#'*62}")
    print(f"  FILE SALVATO: {OUTPUT_PATH}")
    print(f"{'#'*62}\n")

    # Riepilogo match
    print("  RIEPILOGO:")
    for ws in wb.worksheets:
        n = max(ws.max_row - 1, 0) if ws.max_row > 1 else 0
        print(f"    {ws.title:<25} {n} righe")


if __name__ == "__main__":
    main()
