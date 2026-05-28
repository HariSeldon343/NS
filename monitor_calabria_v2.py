#!/usr/bin/env python3
"""
Monitoraggio Atti Calabria - v2
Usa Playwright per portali ASP e TAR, requests per Regione Calabria.
"""

import re
import time
import logging
from datetime import date, timedelta
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)

# ── dates ──────────────────────────────────────────────────────────────────────
TODAY = date(2026, 5, 28)
DATE_FROM_2D = TODAY - timedelta(days=2)  # 2026-05-26
DATE_FROM_3D = TODAY - timedelta(days=3)  # 2026-05-25

# ── keywords ──────────────────────────────────────────────────────────────────
KEYWORDS_NORMAL = [
    "ADI",
    "Assistenza domiciliare",
    "ANMIC",
    "Accreditamento",
    "Aumento di budget",
    "Autismo",
    "Autorizzazione all'esercizio",
    "Autorizzazione alla realizzazione",
    "Autorizzazioni",
    "Budget",
    "Casa Giardino",
    "Centro San Giuseppe",
    "Centro salute e benessere",
    "Fabbisogni LEA",
    "Fisiolab",
    "Fisioterapia",
    "Parere commissione",
    "Presa d'atto verifica",
    "Programmazione",
    "Rete riabilitativa",
    "Rete territoriale",
    "Riabilitazione estensiva",
    "Riconversione prestazioni",
    "Rinnovo accreditamento",
    "San Teodoro",
    "Savelli Hospital",
    "Starbene",
    "Verifica requisiti",
    "Villa San Giuseppe",
    "Villa del Rosario",
]
KEYWORD_LIFE_RE = re.compile(r'\bLIFE\b', re.IGNORECASE)
KEYWORDS_RE = re.compile('|'.join(re.escape(kw) for kw in KEYWORDS_NORMAL), re.IGNORECASE)


def matches_keywords(text: str) -> bool:
    if not text:
        return False
    return bool(KEYWORDS_RE.search(text)) or bool(KEYWORD_LIFE_RE.search(text))


# ── HTTP session (for Regione Calabria) ───────────────────────────────────────
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "it-IT,it;q=0.9,en;q=0.8",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}

SESSION = requests.Session()
SESSION.headers.update(HEADERS)


def safe_get(url, params=None, timeout=30, retries=3, **kwargs):
    for attempt in range(retries):
        try:
            r = SESSION.get(url, params=params, timeout=timeout, **kwargs)
            r.raise_for_status()
            return r
        except Exception as exc:
            log.warning("GET %s attempt %d/%d: %s", url, attempt + 1, retries, exc)
            if attempt < retries - 1:
                time.sleep(2 ** attempt)
    return None


def safe_post(url, data=None, timeout=30, retries=3, **kwargs):
    for attempt in range(retries):
        try:
            r = SESSION.post(url, data=data, timeout=timeout, **kwargs)
            r.raise_for_status()
            return r
        except Exception as exc:
            log.warning("POST %s attempt %d/%d: %s", url, attempt + 1, retries, exc)
            if attempt < retries - 1:
                time.sleep(2 ** attempt)
    return None


# ── ASP portals via Playwright ────────────────────────────────────────────────
ASP_PORTALS = [
    ("ASP Cosenza",         "https://online-aspco.sisr.regione.calabria.it"),
    ("ASP Catanzaro",       "https://online-aspcz.sisr.regione.calabria.it"),
    ("ASP Crotone",         "https://online-aspkr.sisr.regione.calabria.it"),
    ("ASP Reggio Calabria", "https://online-asprc.sisr.regione.calabria.it"),
    ("ASP Vibo Valentia",   "https://online-aspvv.sisr.regione.calabria.it"),
]
SEARCH_PATH = "/AlboOnline/ricercaAlbo"


def fetch_asp_playwright(pw, name: str, base_url: str) -> list[dict]:
    results = []
    search_url = base_url + SEARCH_PATH
    date_from_str = DATE_FROM_2D.strftime("%Y-%m-%d")
    log.info("[%s] Opening %s", name, search_url)

    browser = pw.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
    context = browser.new_context(
        user_agent=HEADERS["User-Agent"],
        locale="it-IT",
        ignore_https_errors=True,
    )
    page = context.new_page()

    try:
        page.goto(search_url, timeout=30000, wait_until="domcontentloaded")
        page.wait_for_timeout(2000)

        # Check if page loaded
        title = page.title()
        log.info("[%s] Page title: %s", name, title)

        # Fill date field
        date_filled = False
        for sel in ["#dataPubblicazioneDal", "input[id='dataPubblicazioneDal']",
                    "input[name='dataPubblicazioneDal']", "input[type='date']"]:
            try:
                el = page.query_selector(sel)
                if el and el.is_visible():
                    el.fill(date_from_str)
                    date_filled = True
                    log.info("[%s] Filled date with selector: %s", name, sel)
                    break
            except Exception:
                pass

        if not date_filled:
            log.warning("[%s] Could not find date field", name)
            # Try to get all input fields
            inputs = page.query_selector_all("input")
            for inp in inputs:
                inp_id = inp.get_attribute("id") or ""
                inp_name = inp.get_attribute("name") or ""
                inp_type = inp.get_attribute("type") or ""
                log.info("[%s]   input id=%s name=%s type=%s", name, inp_id, inp_name, inp_type)

        # Submit form
        submitted = False
        for sel in ["button[type='submit']", "input[type='submit']", ".btn-primary", ".btn-search"]:
            try:
                el = page.query_selector(sel)
                if el and el.is_visible():
                    el.click()
                    page.wait_for_timeout(3000)
                    submitted = True
                    log.info("[%s] Submitted via: %s", name, sel)
                    break
            except Exception:
                pass

        if not submitted:
            page.keyboard.press("Enter")
            page.wait_for_timeout(3000)

        # Iterate pages
        page_num = 1
        while True:
            log.info("[%s] Parsing page %d", name, page_num)
            html = page.content()
            rows = _parse_asp_html(html, base_url)
            log.info("[%s] Page %d: %d rows found", name, page_num, len(rows))

            for row in rows:
                if matches_keywords(row.get("oggetto", "")):
                    results.append(row)

            # Try next page
            next_found = False
            for next_sel in [
                f"a[href*='page={page_num + 1}']",
                f"a:text('{page_num + 1}')",
                "a.next", ".next a", "a[rel='next']",
                f"//a[contains(@class,'page') and text()='{page_num + 1}']",
            ]:
                try:
                    if next_sel.startswith("//"):
                        el = page.query_selector(f"xpath={next_sel}")
                    else:
                        el = page.query_selector(next_sel)
                    if el and el.is_visible():
                        el.click()
                        page.wait_for_timeout(2000)
                        next_found = True
                        page_num += 1
                        break
                except Exception:
                    pass

            if not next_found:
                break
            if page_num > 50:
                log.warning("[%s] Safety break at page 50", name)
                break

    except Exception as exc:
        log.error("[%s] Error: %s", name, exc)
        # Save screenshot for debugging
        try:
            page.screenshot(path=f"/tmp/{name.replace(' ','_')}_error.png")
        except Exception:
            pass
        # Try to get HTML anyway
        try:
            html = page.content()
            log.info("[%s] Got HTML anyway (%d bytes)", name, len(html))
            rows = _parse_asp_html(html, base_url)
            for row in rows:
                if matches_keywords(row.get("oggetto", "")):
                    results.append(row)
        except Exception:
            pass
    finally:
        browser.close()

    log.info("[%s] Total matching: %d", name, len(results))
    return results


def _parse_asp_html(html: str, base_url: str) -> list[dict]:
    rows = []
    soup = BeautifulSoup(html, "lxml")

    table = None
    # Try multiple selectors
    for sel in [
        {"id": re.compile(r"risultati|atti|albo|tabellaRisultati", re.I)},
        {"class": re.compile(r"risultati|atti|albo|table", re.I)},
    ]:
        table = soup.find("table", sel)
        if table:
            break
    if not table:
        table = soup.find("table")
    if not table:
        return rows

    # Get headers
    headers = []
    thead = table.find("thead")
    if thead:
        headers = [th.get_text(strip=True).lower() for th in thead.find_all(["th", "td"])]

    tbody = table.find("tbody") or table
    for tr in tbody.find_all("tr"):
        cells = tr.find_all(["td", "th"])
        if not cells or (len(cells) < 2):
            continue
        if all(c.name == "th" for c in cells):
            continue

        texts = [c.get_text(strip=True) for c in cells]
        row_data = {}
        if headers and len(headers) == len(cells):
            for h, t in zip(headers, texts):
                row_data[h] = t
        else:
            # Guess: try to find date-like and text-like columns
            for i, t in enumerate(texts):
                row_data[f"col{i}"] = t

        # Find link
        link = ""
        for a in tr.find_all("a", href=True):
            href = a["href"]
            if href and href != "#" and not href.startswith("javascript"):
                link = urljoin(base_url, href)
                break

        # Extract oggetto
        oggetto = ""
        for k in ["oggetto", "descrizione", "titolo", "oggetto/descrizione", "col3", "col4", "col2"]:
            if k in row_data and row_data[k]:
                oggetto = row_data[k]
                break
        if not oggetto:
            # Use longest cell text
            longest = max(texts, key=len, default="")
            oggetto = longest

        # Extract data
        data_val = ""
        for k in ["data", "data pubblicazione", "data pubbl.", "datapubblicazione", "col0", "col2"]:
            if k in row_data and row_data[k]:
                data_val = row_data[k]
                break
        if not data_val:
            # Find date-like string
            for t in texts:
                if re.match(r'\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}', t):
                    data_val = t
                    break

        rows.append({"data": data_val, "oggetto": oggetto, "link": link, "_raw": texts})

    return rows


# ── Regione Calabria via requests ─────────────────────────────────────────────
RC_BASE = "https://www.regione.calabria.it"
RC_PATH = "/provvedimenti-della-regione/"


def fetch_regione_calabria() -> list[dict]:
    results = []
    date_from_str = DATE_FROM_2D.strftime("%Y-%m-%d")
    base_url = RC_BASE + RC_PATH
    log.info("[Regione Calabria] Starting, date_from=%s", date_from_str)

    page_num = 1
    while True:
        log.info("[Regione Calabria] Page %d", page_num)

        # Regione Calabria uses GET with paged= and filter params
        params = {
            "filter_active": "true",
            "filter_date_from": date_from_str,
        }
        if page_num > 1:
            params["paged"] = str(page_num)

        r = safe_get(base_url, params=params)
        if r is None:
            log.error("[Regione Calabria] Failed on page %d", page_num)
            break

        soup = BeautifulSoup(r.text, "lxml")
        table = soup.find("table", class_="table")
        if not table:
            log.info("[Regione Calabria] No table on page %d", page_num)
            break

        tbody = table.find("tbody")
        if not tbody:
            break

        tr_list = tbody.find_all("tr")
        if not tr_list:
            log.info("[Regione Calabria] No rows on page %d", page_num)
            break

        log.info("[Regione Calabria] Page %d: %d rows", page_num, len(tr_list))

        for tr in tr_list:
            cells = tr.find_all(["td", "th"])
            if len(cells) < 5:
                continue

            tipologia = cells[0].get_text(strip=True)
            data_rep = cells[1].get_text(strip=True)
            numero = cells[2].get_text(strip=True)
            dipartimento = cells[3].get_text(strip=True)
            oggetto = cells[4].get_text(strip=True)

            link = ""
            a_tag = tr.find("a", href=True)
            if a_tag:
                link = urljoin(RC_BASE, a_tag["href"])

            if matches_keywords(oggetto):
                results.append({
                    "data": data_rep,
                    "oggetto": oggetto,
                    "link": link,
                    "_extra": f"{tipologia} N.{numero} - {dipartimento}",
                })

        # Check for next page
        pag = soup.find("ul", class_="page-numbers")
        if not pag:
            break

        has_next = False
        for a in pag.find_all("a", class_="page-numbers"):
            href = a.get("href", "")
            if f"paged={page_num + 1}" in href:
                has_next = True
                break

        if not has_next:
            break
        page_num += 1

        if page_num > 200:
            log.warning("[Regione Calabria] Safety break at page 200")
            break

    log.info("[Regione Calabria] Total matching: %d", len(results))
    return results


# ── TAR Catanzaro via Playwright ──────────────────────────────────────────────
TAR_BASE = "https://www.giustizia-amministrativa.it"
TAR_PATH = "/provvedimenti-tar-catanzaro"


def fetch_tar_playwright(pw) -> list[dict]:
    results = []
    date_from_str = DATE_FROM_3D.strftime("%Y-%m-%d")
    url = TAR_BASE + TAR_PATH
    log.info("[TAR Catanzaro] Opening %s", url)

    browser = pw.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
    context = browser.new_context(
        user_agent=HEADERS["User-Agent"],
        locale="it-IT",
        ignore_https_errors=True,
    )
    page = context.new_page()

    try:
        page.goto(url, timeout=30000, wait_until="domcontentloaded")
        page.wait_for_timeout(3000)
        log.info("[TAR Catanzaro] Title: %s", page.title())

        # Try to set date filter
        date_set = False
        for sel in [
            "input[name='publishDateFrom']",
            "#publishDateFrom",
            "input[id*='dateFrom']",
            "input[type='date']",
        ]:
            try:
                el = page.query_selector(sel)
                if el and el.is_visible():
                    el.fill(date_from_str)
                    date_set = True
                    log.info("[TAR Catanzaro] Set date via: %s", sel)
                    break
            except Exception:
                pass

        if not date_set:
            log.warning("[TAR Catanzaro] Could not set date filter")
            # Log all inputs
            inputs = page.query_selector_all("input")
            for inp in inputs:
                log.info("[TAR]   input id=%s name=%s type=%s",
                         inp.get_attribute("id"), inp.get_attribute("name"), inp.get_attribute("type"))

        # Submit
        for sel in ["button[type='submit']", "input[type='submit']", ".btn-primary", ".btn-search", ".search-btn"]:
            try:
                el = page.query_selector(sel)
                if el and el.is_visible():
                    el.click()
                    page.wait_for_timeout(3000)
                    log.info("[TAR Catanzaro] Clicked submit: %s", sel)
                    break
            except Exception:
                pass

        page_num = 1
        while True:
            log.info("[TAR Catanzaro] Parsing page %d", page_num)
            html = page.content()
            rows = _parse_tar_html(html)
            log.info("[TAR Catanzaro] Page %d: %d rows", page_num, len(rows))

            for row in rows:
                search_text = (row.get("parte", "") + " " + row.get("oggetto", "")).strip()
                if matches_keywords(search_text):
                    results.append(row)

            # Next page
            next_found = False
            for next_sel in [
                f"a:text('{page_num + 1}')",
                "a.next", ".next a", "a[rel='next']",
                ".pagination a.active + a",
            ]:
                try:
                    el = page.query_selector(next_sel)
                    if el and el.is_visible():
                        el.click()
                        page.wait_for_timeout(2000)
                        next_found = True
                        page_num += 1
                        break
                except Exception:
                    pass

            if not next_found:
                break
            if page_num > 50:
                break

    except Exception as exc:
        log.error("[TAR Catanzaro] Error: %s", exc)
        try:
            html = page.content()
            rows = _parse_tar_html(html)
            for row in rows:
                search_text = (row.get("parte", "") + " " + row.get("oggetto", "")).strip()
                if matches_keywords(search_text):
                    results.append(row)
        except Exception:
            pass
    finally:
        browser.close()

    log.info("[TAR Catanzaro] Total matching: %d", len(results))
    return results


def _parse_tar_html(html: str) -> list[dict]:
    rows = []
    soup = BeautifulSoup(html, "lxml")
    table = soup.find("table")
    if not table:
        return rows

    headers = []
    thead = table.find("thead")
    if thead:
        headers = [th.get_text(strip=True).lower() for th in thead.find_all(["th", "td"])]

    tbody = table.find("tbody") or table
    for tr in tbody.find_all("tr"):
        cells = tr.find_all(["td", "th"])
        if not cells or all(c.name == "th" for c in cells):
            continue

        texts = [c.get_text(strip=True) for c in cells]
        row_data = {}
        if headers and len(headers) == len(cells):
            for h, t in zip(headers, texts):
                row_data[h] = t
        else:
            for i, t in enumerate(texts):
                row_data[f"col{i}"] = t

        link = ""
        for a in tr.find_all("a", href=True):
            href = a["href"]
            if href and href != "#":
                link = urljoin(TAR_BASE, href)
                break

        parte = row_data.get("parte") or row_data.get("ricorrente") or row_data.get("col2") or ""
        oggetto = row_data.get("oggetto") or row_data.get("descrizione") or row_data.get("col3") or ""
        data_val = row_data.get("data") or row_data.get("data pubblicazione") or row_data.get("col0") or ""

        rows.append({"data": data_val, "parte": parte, "oggetto": oggetto, "link": link, "_raw": texts})

    return rows


# ── Excel builder ─────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
LINK_FONT = Font(color="0563C1", underline="single")
ALT_FILL = PatternFill("solid", fgColor="DCE6F1")
COLUMNS = ["Data", "Oggetto", "Descrizione breve", "Link"]
COL_WIDTHS = [15, 65, 55, 12]


def build_excel(data_by_sheet: dict, out_path: str):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sheet_order = [
        "ASP Cosenza",
        "ASP Catanzaro",
        "ASP Crotone",
        "ASP Reggio Calabria",
        "ASP Vibo Valentia",
        "Regione Calabria",
        "TAR Catanzaro",
    ]

    for sheet_name in sheet_order:
        rows = data_by_sheet.get(sheet_name, [])
        ws = wb.create_sheet(title=sheet_name[:31])

        for col_idx, (col_name, width) in enumerate(zip(COLUMNS, COL_WIDTHS), start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.row_dimensions[1].height = 22

        if not rows:
            ws.cell(row=2, column=1, value="Nessun risultato")
            continue

        for row_idx, row in enumerate(rows, start=2):
            oggetto = row.get("oggetto") or ""
            if not oggetto and "parte" in row:
                oggetto = row["parte"]
            data_val = row.get("data", "")
            link_url = row.get("link", "")
            descrizione = (oggetto[:150]) if oggetto else ""

            fill = ALT_FILL if row_idx % 2 == 0 else None

            c_data = ws.cell(row=row_idx, column=1, value=data_val)
            c_data.alignment = Alignment(vertical="top")
            if fill:
                c_data.fill = fill

            c_ogg = ws.cell(row=row_idx, column=2, value=oggetto)
            c_ogg.alignment = Alignment(wrap_text=True, vertical="top")
            if fill:
                c_ogg.fill = fill

            c_desc = ws.cell(row=row_idx, column=3, value=descrizione)
            c_desc.alignment = Alignment(wrap_text=True, vertical="top")
            if fill:
                c_desc.fill = fill

            c_link = ws.cell(row=row_idx, column=4, value="Apri" if link_url else "—")
            if link_url:
                c_link.hyperlink = link_url
                c_link.font = LINK_FONT
            c_link.alignment = Alignment(vertical="top", horizontal="center")
            if fill:
                c_link.fill = fill

        ws.freeze_panes = "A2"

    wb.save(out_path)
    log.info("Saved: %s", out_path)


# ── main ──────────────────────────────────────────────────────────────────────
def main():
    data_by_sheet = {}

    with sync_playwright() as pw:
        # ASP portals
        for name, base_url in ASP_PORTALS:
            data_by_sheet[name] = fetch_asp_playwright(pw, name, base_url)

        # TAR Catanzaro
        data_by_sheet["TAR Catanzaro"] = fetch_tar_playwright(pw)

    # Regione Calabria (no browser needed)
    data_by_sheet["Regione Calabria"] = fetch_regione_calabria()

    # Output file
    date_str = TODAY.strftime("%d-%m-%Y")
    filename = f"Monitoraggio_Atti_Calabria_{date_str}.xlsx"
    out_path = f"/home/user/NS/{filename}"

    build_excel(data_by_sheet, out_path)

    print("\n" + "=" * 60)
    print(f"File: {filename}")
    print("=" * 60)
    for sheet_name in [
        "ASP Cosenza", "ASP Catanzaro", "ASP Crotone",
        "ASP Reggio Calabria", "ASP Vibo Valentia",
        "Regione Calabria", "TAR Catanzaro",
    ]:
        n = len(data_by_sheet.get(sheet_name, []))
        print(f"  {sheet_name:30s}: {n} atti trovati")
    print("=" * 60)

    return out_path


if __name__ == "__main__":
    main()
