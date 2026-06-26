"""
Monitoraggio Atti Calabria - Scraper completo
Portali: 5x ASP Calabria, Regione Calabria, TAR Catanzaro
"""

import requests
import json
import re
import datetime
import sys
import time
from urllib.parse import urlencode
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# KEYWORDS
# ---------------------------------------------------------------------------
KEYWORDS_PLAIN = [
    "ADI",
    "Assistenza domiciliare",
    "ANMIC",
    "Accreditamento",
    "Aumento di budget",
    "Autismo",
    "Autorizzazione all'esercizio",
    "Autorizzazione alla realizzazione",
    "Autorizzazioni",
    "Budget",
    "Casa Giardino",
    "Centro San Giuseppe",
    "Centro salute e benessere",
    "Fabbisogni LEA",
    "Fisiolab",
    "Fisioterapia",
    "Parere commissione",
    "Presa d'atto verifica",
    "Programmazione",
    "Rete riabilitativa",
    "Rete territoriale",
    "Riabilitazione estensiva",
    "Riconversione prestazioni",
    "Rinnovo accreditamento",
    "San Teodoro",
    "Savelli Hospital",
    "Starbene",
    "Verifica requisiti",
    "Villa San Giuseppe",
    "Villa del Rosario",
]

def matches_keywords(text: str) -> bool:
    if not text:
        return False
    t = text.lower()
    for kw in KEYWORDS_PLAIN:
        if kw.lower() in t:
            return True
    # "Life" con word boundary
    if re.search(r'\blife\b', t):
        return True
    return False

# ---------------------------------------------------------------------------
# SESSION
# ---------------------------------------------------------------------------
SESSION = requests.Session()
SESSION.verify = '/root/.ccr/ca-bundle.crt'
SESSION.headers.update({
    'User-Agent': (
        'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
        'AppleWebKit/537.36 (KHTML, like Gecko) '
        'Chrome/120.0.0.0 Safari/537.36'
    )
})

TODAY = datetime.date.today()
TWO_DAYS_AGO  = TODAY - datetime.timedelta(days=2)
THREE_DAYS_AGO = TODAY - datetime.timedelta(days=3)

# ---------------------------------------------------------------------------
# ASP portals (5 portali – inaccessibili dalla rete corrente)
# ---------------------------------------------------------------------------
ASP_PORTALS = {
    'ASP Cosenza':         'https://online-aspco.sisr.regione.calabria.it/AlboOnline/ricercaAlbo',
    'ASP Catanzaro':       'https://online-aspcz.sisr.regione.calabria.it/AlboOnline/ricercaAlbo',
    'ASP Crotone':         'https://online-aspkr.sisr.regione.calabria.it/AlboOnline/ricercaAlbo',
    'ASP Reggio Calabria': 'https://online-asprc.sisr.regione.calabria.it/AlboOnline/ricercaAlbo',
    'ASP Vibo Valentia':   'https://online-aspvv.sisr.regione.calabria.it/AlboOnline/ricercaAlbo',
}

def try_asp_portal(name: str, base_url: str, date_from: datetime.date, date_to: datetime.date):
    """Tenta di raschiare un portale ASP Calabria (ricercaAlbo)."""
    print(f"  Tentativo connessione: {name} ...")
    results = []
    error_msg = None

    try:
        # GET iniziale per ottenere la struttura del form
        r = SESSION.get(base_url, timeout=15)
        r.raise_for_status()
        soup = BeautifulSoup(r.text, 'lxml')

        # Il form tipico ha id="dataPubblicazioneDal" per la data
        form = soup.find('form')
        if not form:
            error_msg = "Struttura form non rilevata"
            return [], error_msg

        # Costruiamo i dati POST
        # I campi standard di questi portali Regione Calabria (SISR)
        data_from_str = date_from.strftime('%Y-%m-%d')
        data_to_str   = date_to.strftime('%Y-%m-%d')

        form_data = {
            'dataPubblicazioneDal': data_from_str,
            'dataPubblicazioneAl':  data_to_str,
        }

        # Raccoglie hidden inputs
        for inp in form.find_all('input', {'type': 'hidden'}):
            if inp.get('name') and inp.get('value') is not None:
                form_data[inp['name']] = inp['value']

        action = form.get('action', base_url)
        if not action.startswith('http'):
            action = base_url.rsplit('/', 1)[0] + '/' + action.lstrip('/')

        page = 1
        while True:
            form_data['page'] = str(page)
            resp = SESSION.post(action, data=form_data, timeout=20)
            soup2 = BeautifulSoup(resp.text, 'lxml')

            rows = soup2.select('table tbody tr')
            if not rows:
                break

            found_any = False
            for row in rows:
                cells = row.find_all(['td', 'th'])
                texts = [c.get_text(separator=' ', strip=True) for c in cells]
                if not texts:
                    continue
                # Oggetto è tipicamente nella penultima o terza colonna
                oggetto = ' '.join(texts)
                link_tag = row.find('a', href=True)
                link = link_tag['href'] if link_tag else ''
                if link and not link.startswith('http'):
                    link = base_url.rsplit('/', 1)[0] + '/' + link.lstrip('/')

                # Trova la cella con la data (formato DD/MM/YYYY o YYYY-MM-DD)
                data_cell = ''
                for t in texts:
                    if re.search(r'\d{2}/\d{2}/\d{4}|\d{4}-\d{2}-\d{2}', t):
                        data_cell = t
                        break

                if matches_keywords(oggetto):
                    results.append({
                        'data':    data_cell,
                        'oggetto': oggetto,
                        'link':    link,
                    })
                found_any = True

            # Pagina successiva
            next_btn = soup2.find('a', string=re.compile(r'successiv|next|>>', re.I))
            if not next_btn or not found_any:
                break
            page += 1

    except requests.exceptions.SSLError:
        error_msg = (
            "Portale non raggiungibile: errore TLS/SSL durante la connessione. "
            "I portali sisr.regione.calabria.it sono probabilmente accessibili solo da IP italiani "
            "o utilizzano configurazioni TLS non compatibili con questa rete. "
            f"URL: {base_url}"
        )
    except requests.exceptions.ConnectionError:
        error_msg = (
            "Portale non raggiungibile: connessione rifiutata o timeout. "
            f"URL: {base_url}"
        )
    except Exception as e:
        error_msg = f"ERRORE imprevisto: {type(e).__name__}: {e}"

    if error_msg:
        print(f"    {error_msg}")

    return results, error_msg


# ---------------------------------------------------------------------------
# REGIONE CALABRIA
# ---------------------------------------------------------------------------
def scrape_regione_calabria(date_from: datetime.date, date_to: datetime.date):
    print("  Scraping Regione Calabria ...")
    results = []
    base = "https://www.regione.calabria.it/provvedimenti-della-regione/"
    paged = 1

    while True:
        params = {
            'filter_date_from': date_from.strftime('%Y-%m-%d'),
            'filter_date_to':   date_to.strftime('%Y-%m-%d'),
            'sort_order': '0',
            'pageNum':    '0',
            'paged':      str(paged),
        }
        url = base + '?' + urlencode(params)
        print(f"    Pagina {paged}: {url}")

        try:
            r = SESSION.get(url, timeout=30)
            r.raise_for_status()
        except Exception as e:
            print(f"    Errore pagina {paged}: {e}")
            break

        soup = BeautifulSoup(r.text, 'lxml')
        tbody = soup.find('tbody')
        if not tbody:
            break

        rows = tbody.find_all('tr')
        if not rows:
            break

        page_had_rows = False
        for row in rows:
            cells = row.find_all(['td', 'th'])
            if len(cells) < 5:
                continue
            page_had_rows = True

            # Struttura: Tipologia | Data Repertoriazione | N | Dipartimento | Oggetto | Dettaglio
            data_text   = cells[1].get_text(strip=True)  # Data Repertoriazione
            oggetto     = cells[4].get_text(strip=True)  # Oggetto
            link_cell   = cells[5] if len(cells) > 5 else None
            link        = ''
            if link_cell:
                a = link_cell.find('a', href=True)
                if a:
                    href = a['href']
                    link = href if href.startswith('http') else 'https://www.regione.calabria.it' + href

            if matches_keywords(oggetto):
                results.append({
                    'data':    data_text,
                    'oggetto': oggetto,
                    'link':    link,
                })

        # Controlla se c'è pagina successiva
        next_a = soup.find('a', class_='next page-numbers')
        if not next_a or not page_had_rows:
            break
        paged += 1
        time.sleep(0.3)

    print(f"    Trovati {len(results)} risultati con match keyword")
    return results, None


# ---------------------------------------------------------------------------
# TAR CATANZARO
# ---------------------------------------------------------------------------
PORTLET_ID = (
    "it_indra_ga_institutional_area_"
    "JurisdictionalActivityAdministrativeActsWebPortlet_INSTANCE_jjYpzZYF4Qfe"
)
NS = f"_{PORTLET_ID}_"


def build_tar_payload(date_from: datetime.date, date_to: datetime.date,
                      start: int = 0, length: int = 500) -> dict:
    additional_info = {
        "schema":              "TAR_CATANZARO",
        "type":                None,
        "year":                "",
        "number":              "",
        "hearingDateFrom":     None,
        "hearingDateTo":       None,
        "publishDateFrom":     date_from.strftime('%Y-%m-%d'),
        "publishDateTo":       date_to.strftime('%Y-%m-%d'),
        "hearingType":         None,
        "nrg":                 None,
        "section":             "",
        "provisionSpecification": "",
        "president":           "",
        "draftingJudge":       "",
        "subjectMatter":       None,
        "page":                None,
        "size":                None,
        "orderBy":             None,
        "orderStrategy":       None,
        "queryString":         None,
    }
    cols = [
        "nrgFascicolo", "sezione", "parte", "tipoUdienza", "dataUdienza",
        "numProvvedimento", "dataPubblicazione", "tipoProvvedimento",
        "relatore", "presidente", "esito",
    ]
    return {
        "draw":    1,
        "start":   start,
        "length":  length,
        "search":  {"value": "", "regex": False},
        "order":   [{"column": 6, "dir": "desc"}, {"column": 5, "dir": "desc"}],
        "columns": [
            {
                "data": c, "name": "", "searchable": True, "orderable": True,
                "search": {"value": "", "regex": False}
            } for c in cols
        ],
        "additionalInfo": json.dumps(additional_info),
    }


def scrape_tar_catanzaro(date_from: datetime.date, date_to: datetime.date):
    print("  Scraping TAR Catanzaro ...")
    results = []

    # Step 1: Ottieni cookies + p_auth + formDate
    try:
        init_r = SESSION.get(
            'https://www.giustizia-amministrativa.it/provvedimenti-tar-catanzaro',
            timeout=30
        )
        init_r.raise_for_status()
    except Exception as e:
        return [], f"ERRORE connessione TAR: {e}"

    p_auth_m = re.search(r'p_auth=([A-Za-z0-9]+)', init_r.text)
    p_auth   = p_auth_m.group(1) if p_auth_m else ''

    form_date_m = re.search(
        rf'name="{re.escape(NS)}formDate"[^>]+value="(\d+)"',
        init_r.text
    )
    form_date = form_date_m.group(1) if form_date_m else str(int(datetime.datetime.now().timestamp() * 1000))

    # Step 2: POST lifecycle=1 per stabilire lo stato della ricerca
    action_url = (
        f"https://www.giustizia-amministrativa.it/web/guest/provvedimenti-tar-catanzaro"
        f"?p_p_id={PORTLET_ID}&p_p_lifecycle=1&p_p_state=normal&p_p_mode=view"
        f"&{NS}javax.portlet.action=%2Fadministrative-acts%2Fsearch&p_auth={p_auth}"
    )
    post_data = {
        f'{NS}formDate':        form_date,
        f'{NS}year':            '',
        f'{NS}number':          '',
        f'{NS}section':         '',
        f'{NS}type':            '',
        f'{NS}specific':        '',
        f'{NS}president':       '',
        f'{NS}draftingJudge':   '',
        f'{NS}hearingDateFrom': '',
        f'{NS}hearingDateTo':   '',
        f'{NS}publishDateFrom': date_from.strftime('%Y-%m-%d'),
        f'{NS}publishDateTo':   date_to.strftime('%Y-%m-%d'),
    }
    post_hdrs = {
        'Referer':         'https://www.giustizia-amministrativa.it/provvedimenti-tar-catanzaro',
        'X-Requested-With': 'XMLHttpRequest',
        'Content-Type':    'application/x-www-form-urlencoded',
    }
    try:
        SESSION.post(action_url, data=post_data, headers=post_hdrs, timeout=30)
    except Exception as e:
        print(f"    Warning lifecycle=1: {e}")

    # Step 3: Call DataTables AJAX resource (lifecycle=2)
    resource_url = (
        f"https://www.giustizia-amministrativa.it/web/guest/provvedimenti-tar-catanzaro"
        f"?p_p_id={PORTLET_ID}&p_p_lifecycle=2&p_p_state=normal&p_p_mode=view"
        f"&p_p_resource_id=%2Fadministrative-acts%2Fsearch%2Fresults&p_p_cacheability=cacheLevelPage"
    )
    ajax_hdrs = {
        'Content-Type':     'application/json',
        'Accept':           'application/json',
        'X-Requested-With': 'XMLHttpRequest',
        'Referer':          'https://www.giustizia-amministrativa.it/provvedimenti-tar-catanzaro',
    }

    start = 0
    length = 500
    while True:
        payload = build_tar_payload(date_from, date_to, start=start, length=length)
        try:
            r = SESSION.post(resource_url, json=payload, headers=ajax_hdrs, timeout=30)
            r.raise_for_status()
            data = r.json()
        except Exception as e:
            return results, f"ERRORE fetch TAR dati: {e}"

        records       = data.get('data', [])
        records_total = data.get('recordsTotal', 0)
        print(f"    Batch start={start}: {len(records)} record (totale {records_total})")

        for rec in records:
            parte     = rec.get('parte', '') or ''
            nrg       = rec.get('nrgFascicolo', '') or ''
            nome_file = rec.get('nomeFile', '') or ''
            data_pub  = rec.get('dataPubblicazione', '') or ''
            tipo      = rec.get('tipoProvvedimento', '') or ''
            num_prov  = rec.get('numProvvedimento', '') or ''

            # Per TAR filtriamo su "Parte"
            if matches_keywords(parte):
                if nome_file:
                    link = (
                        f"https://mdp.giustizia-amministrativa.it/visualizza/"
                        f"?nodeRef=&schema=tar_cz&nrg={nrg}&nomeFile={nome_file}&subDir=Provvedimenti"
                    )
                else:
                    link = ''

                oggetto_display = (
                    f"{tipo} n.{num_prov} - {parte}"
                    if tipo and num_prov else parte
                )

                results.append({
                    'data':    data_pub,
                    'oggetto': oggetto_display,
                    'link':    link,
                })

        if start + length >= records_total:
            break
        start += length
        time.sleep(0.2)

    print(f"    Trovati {len(results)} risultati TAR con match keyword")
    return results, None


# ---------------------------------------------------------------------------
# EXCEL WRITER
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
LINK_FONT   = Font(color="0563C1", underline="single")

def write_sheet(ws, results, error_msg=None):
    """Scrive i risultati in un foglio Excel."""
    headers = ["Data", "Oggetto", "Descrizione breve (150 car.)", "Link atto"]

    # Header row
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font   = HEADER_FONT
        cell.fill   = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True)

    # Larghezze colonne
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 50

    ws.freeze_panes = "A2"

    if error_msg:
        ws.cell(row=2, column=1, value=f"ERRORE: {error_msg}")
        return

    if not results:
        ws.cell(row=2, column=1, value="Nessun risultato")
        return

    for row_idx, rec in enumerate(results, 2):
        data_val   = rec.get('data', '')
        oggetto    = rec.get('oggetto', '')
        link       = rec.get('link', '')
        desc_breve = oggetto[:150] if oggetto else ''

        ws.cell(row=row_idx, column=1, value=data_val)
        ws.cell(row=row_idx, column=2, value=oggetto).alignment = Alignment(wrap_text=True)
        ws.cell(row=row_idx, column=3, value=desc_breve)

        if link:
            link_cell = ws.cell(row=row_idx, column=4, value="Apri atto")
            link_cell.hyperlink = link
            link_cell.font      = LINK_FONT
        else:
            ws.cell(row=row_idx, column=4, value="N/D")


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------
def main():
    today_str = TODAY.strftime('%d-%m-%Y')
    filename  = f"Monitoraggio_Atti_Calabria_{today_str}.xlsx"
    outpath   = f"/tmp/claude-0/-home-user-NS/8200bed8-442f-5e1c-8fca-af13c620f30b/scratchpad/{filename}"

    wb = Workbook()
    wb.remove(wb.active)  # rimuove foglio vuoto di default

    # ------------------------------------------------------------------
    # 1-5: Portali ASP (tentiamo, gestiamo errori di connessione)
    # ------------------------------------------------------------------
    for sheet_name, asp_url in ASP_PORTALS.items():
        print(f"\n[{sheet_name}]")
        ws = wb.create_sheet(title=sheet_name)
        results, error = try_asp_portal(sheet_name, asp_url, TWO_DAYS_AGO, TODAY)
        write_sheet(ws, results, error)

    # ------------------------------------------------------------------
    # 6: Regione Calabria
    # ------------------------------------------------------------------
    print("\n[Regione Calabria]")
    ws_rc = wb.create_sheet(title="Regione Calabria")
    rc_results, rc_error = scrape_regione_calabria(TWO_DAYS_AGO, TODAY)
    write_sheet(ws_rc, rc_results, rc_error)

    # ------------------------------------------------------------------
    # 7: TAR Catanzaro
    # ------------------------------------------------------------------
    print("\n[TAR Catanzaro]")
    ws_tar = wb.create_sheet(title="TAR Catanzaro")
    tar_results, tar_error = scrape_tar_catanzaro(THREE_DAYS_AGO, TODAY)
    write_sheet(ws_tar, tar_results, tar_error)

    # ------------------------------------------------------------------
    # Salva
    # ------------------------------------------------------------------
    wb.save(outpath)
    print(f"\n✓ File salvato: {outpath}")
    print(f"  Regione Calabria: {len(rc_results)} atti matchati")
    print(f"  TAR Catanzaro:    {len(tar_results)} atti matchati")
    return outpath


if __name__ == '__main__':
    out = main()
    print(f"\nPATH: {out}")
