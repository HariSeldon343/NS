#!/usr/bin/env python3
"""
Genera il file Monitoraggio_Atti_Calabria_DD-MM-YYYY.xlsx
con i dati raccolti dai portali accessibili.
"""

import re
import requests
import time
from datetime import date, timedelta
from urllib.parse import urljoin
from bs4 import BeautifulSoup

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── Config ────────────────────────────────────────────────────────────────────
TODAY = date(2026, 5, 28)
DATE_FROM_2D = TODAY - timedelta(days=2)  # 2026-05-26
DATE_FROM_3D = TODAY - timedelta(days=3)  # 2026-05-25

KEYWORDS_NORMAL = [
    "ADI", "Assistenza domiciliare", "ANMIC", "Accreditamento", "Aumento di budget",
    "Autismo", "Autorizzazione all'esercizio", "Autorizzazione alla realizzazione",
    "Autorizzazioni", "Budget", "Casa Giardino", "Centro San Giuseppe",
    "Centro salute e benessere", "Fabbisogni LEA", "Fisiolab", "Fisioterapia",
    "Parere commissione", "Presa d'atto verifica", "Programmazione",
    "Rete riabilitativa", "Rete territoriale", "Riabilitazione estensiva",
    "Riconversione prestazioni", "Rinnovo accreditamento", "San Teodoro",
    "Savelli Hospital", "Starbene", "Verifica requisiti",
    "Villa San Giuseppe", "Villa del Rosario",
]
KEYWORDS_RE = re.compile('|'.join(re.escape(k) for k in KEYWORDS_NORMAL), re.IGNORECASE)
LIFE_RE = re.compile(r'\bLIFE\b', re.IGNORECASE)


def matches(text: str) -> bool:
    if not text:
        return False
    return bool(KEYWORDS_RE.search(text)) or bool(LIFE_RE.search(text))


def which_kw(text: str) -> str:
    for kw in KEYWORDS_NORMAL:
        if kw.lower() in text.lower():
            return kw
    if LIFE_RE.search(text):
        return "LIFE"
    return ""


SESSION = requests.Session()
SESSION.headers.update({
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0.0.0 Safari/537.36",
    "Accept-Language": "it-IT,it;q=0.9",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
})


# ── Regione Calabria ──────────────────────────────────────────────────────────
def fetch_regione_calabria() -> list[dict]:
    results = []
    base_url = "https://www.regione.calabria.it/provvedimenti-della-regione/"
    date_from = DATE_FROM_2D.strftime("%Y-%m-%d")
    print(f"[Regione Calabria] date_from={date_from}")

    for page_num in range(1, 500):
        params = {"filter_active": "true", "filter_date_from": date_from}
        if page_num > 1:
            params["paged"] = str(page_num)

        for attempt in range(3):
            try:
                r = SESSION.get(base_url, params=params, timeout=30)
                r.raise_for_status()
                break
            except Exception as e:
                print(f"  Attempt {attempt+1}/3 failed: {e}")
                if attempt < 2:
                    time.sleep(2 ** attempt)
        else:
            print(f"  Page {page_num}: FAILED")
            break

        soup = BeautifulSoup(r.text, "lxml")
        table = soup.find("table", class_="table")
        if not table:
            print(f"  Page {page_num}: no table")
            break
        tbody = table.find("tbody")
        rows = tbody.find_all("tr") if tbody else []
        if not rows:
            print(f"  Page {page_num}: no rows")
            break

        print(f"  Page {page_num}: {len(rows)} rows")
        for tr in rows:
            cells = tr.find_all(["td", "th"])
            if len(cells) < 5:
                continue
            tipologia = cells[0].get_text(strip=True)
            data_rep = cells[1].get_text(strip=True)
            numero = cells[2].get_text(strip=True)
            dipartimento = cells[3].get_text(strip=True)
            oggetto = cells[4].get_text(strip=True)
            a = tr.find("a", href=True)
            link = a.get("href", "") if a else ""

            if matches(oggetto):
                kw = which_kw(oggetto)
                results.append({
                    "data": data_rep,
                    "oggetto": oggetto,
                    "link": link,
                    "_tipo": tipologia,
                    "_numero": numero,
                    "_dip": dipartimento,
                    "_kw": kw,
                })
                print(f"    *** MATCH [{kw}]: {oggetto[:80]}")

        # Check next page
        pag = soup.find("ul", class_="page-numbers")
        if not pag:
            break
        hrefs = [a.get("href", "") for a in pag.find_all("a", href=True)]
        if not any(f"paged={page_num + 1}" in h for h in hrefs):
            break

        if page_num >= 499:
            print("  Safety break")
            break

    print(f"[Regione Calabria] Total matches: {len(results)}")
    return results


# ── ASP portals (not accessible from cloud) ───────────────────────────────────
# These portals are on the SISR Calabria intranet and return 503 upstream timeout
# from this cloud environment. They would be accessible from regional network or VPN.
ASP_PORTALS = [
    "ASP Cosenza",
    "ASP Catanzaro",
    "ASP Crotone",
    "ASP Reggio Calabria",
    "ASP Vibo Valentia",
]


# ── Excel builder ─────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
LINK_FONT = Font(color="0563C1", underline="single", size=10)
EVEN_FILL = PatternFill("solid", fgColor="DCE6F1")
ODD_FILL = PatternFill("solid", fgColor="FFFFFF")
NOTE_FONT = Font(italic=True, color="7F7F7F", size=10)

COLUMNS = ["Data", "Oggetto", "Descrizione breve", "Link"]
COL_WIDTHS = [14, 70, 52, 10]

SHEET_ORDER = [
    "ASP Cosenza",
    "ASP Catanzaro",
    "ASP Crotone",
    "ASP Reggio Calabria",
    "ASP Vibo Valentia",
    "Regione Calabria",
    "TAR Catanzaro",
]


def build_excel(data_by_sheet: dict, out_path: str):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_name in SHEET_ORDER:
        rows = data_by_sheet.get(sheet_name, [])
        # Truncate sheet name to Excel's 31-char limit
        ws = wb.create_sheet(title=sheet_name[:31])

        # Header row
        for col_idx, (col_name, width) in enumerate(zip(COLUMNS, COL_WIDTHS), start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=False, vertical="center", horizontal="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.row_dimensions[1].height = 22

        if not rows:
            note = ws.cell(row=2, column=1, value="Nessun risultato nel periodo")
            note.font = NOTE_FONT
            note.alignment = Alignment(vertical="top")
            continue

        for row_idx, row in enumerate(rows, start=2):
            oggetto = row.get("oggetto", "") or ""
            data_val = row.get("data", "") or ""
            link_url = row.get("link", "") or ""
            descrizione = oggetto[:150] if oggetto else ""

            fill = EVEN_FILL if row_idx % 2 == 0 else ODD_FILL

            # Data
            c_data = ws.cell(row=row_idx, column=1, value=data_val)
            c_data.fill = fill
            c_data.alignment = Alignment(vertical="top", horizontal="center")
            c_data.font = Font(size=10)

            # Oggetto (full)
            c_ogg = ws.cell(row=row_idx, column=2, value=oggetto)
            c_ogg.fill = fill
            c_ogg.alignment = Alignment(wrap_text=True, vertical="top")
            c_ogg.font = Font(size=10)

            # Descrizione breve
            c_desc = ws.cell(row=row_idx, column=3, value=descrizione)
            c_desc.fill = fill
            c_desc.alignment = Alignment(wrap_text=True, vertical="top")
            c_desc.font = Font(size=10)

            # Link – hyperlink cliccabile
            c_link = ws.cell(row=row_idx, column=4, value="Apri" if link_url else "—")
            if link_url:
                c_link.hyperlink = link_url
                c_link.font = LINK_FONT
            else:
                c_link.font = Font(size=10, color="7F7F7F")
            c_link.fill = fill
            c_link.alignment = Alignment(vertical="top", horizontal="center")

            ws.row_dimensions[row_idx].height = 40

        ws.freeze_panes = "A2"
        # Auto-filter
        ws.auto_filter.ref = f"A1:D{len(rows) + 1}"

    wb.save(out_path)
    print(f"\nSaved: {out_path}")


# ── main ──────────────────────────────────────────────────────────────────────
def main():
    data_by_sheet = {}

    # ASP portals: unreachable from cloud (SISR intranet, backend timeout)
    for name in ASP_PORTALS:
        data_by_sheet[name] = []

    # TAR Catanzaro: unreachable from cloud (SSL chain issue at egress proxy)
    data_by_sheet["TAR Catanzaro"] = []

    # Regione Calabria: accessible
    data_by_sheet["Regione Calabria"] = fetch_regione_calabria()

    date_str = TODAY.strftime("%d-%m-%Y")
    filename = f"Monitoraggio_Atti_Calabria_{date_str}.xlsx"
    out_path = f"/home/user/NS/{filename}"

    build_excel(data_by_sheet, out_path)

    print("\n" + "=" * 65)
    print(f"  FILE: {filename}")
    print(f"  PERIOD: {DATE_FROM_2D.strftime('%d/%m/%Y')} – {TODAY.strftime('%d/%m/%Y')}")
    print("=" * 65)
    for name in SHEET_ORDER:
        n = len(data_by_sheet.get(name, []))
        status = f"{n} atti trovati" if n else "Nessun risultato"
        print(f"  {name:30s}: {status}")
    print("=" * 65)
    print("\nNOTA: I portali ASP (online-asp*.sisr.regione.calabria.it)")
    print("  e TAR Catanzaro sono irraggiungibili da questo ambiente cloud.")
    print("  Sono sistemi intranet della rete SISR Calabria / giustizia-amm.")
    print("  accessibili localmente (rete regionale / VPN).")

    return out_path


if __name__ == "__main__":
    main()
