#!/usr/bin/env python3
"""
Monitoraggio Atti Calabria
Scrapes ASP portals, Regione Calabria, TAR Catanzaro and generates XLSX report.
"""

import requests
from bs4 import BeautifulSoup
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import re
import time
import traceback
from urllib.parse import urljoin
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ─── Dates ────────────────────────────────────────────────────────────────────
TODAY = datetime.now()
DATE_FROM_2 = (TODAY - timedelta(days=2)).strftime('%Y-%m-%d')
DATE_FROM_3 = (TODAY - timedelta(days=3)).strftime('%Y-%m-%d')
DATE_TO     = TODAY.strftime('%Y-%m-%d')
FILE_DATE   = TODAY.strftime('%d-%m-%Y')

# ─── Keywords ────────────────────────────────────────────────────────────────
KEYWORDS = [
    'ADI', 'Assistenza domiciliare', 'ANMIC', 'Accreditamento',
    'Aumento di budget', 'Autismo', "Autorizzazione all'esercizio",
    "Autorizzazione alla realizzazione", 'Autorizzazioni', 'Budget',
    'Casa Giardino', 'Centro San Giuseppe', 'Centro salute e benessere',
    'Fabbisogni LEA', 'Fisiolab', 'Fisioterapia',
    'Parere commissione', "Presa d'atto verifica", 'Programmazione',
    'Rete riabilitativa', 'Rete territoriale', 'Riabilitazione estensiva',
    'Riconversione prestazioni', 'Rinnovo accreditamento',
    'San Teodoro', 'Savelli Hospital', 'Starbene',
    'Verifica requisiti', 'Villa San Giuseppe', 'Villa del Rosario',
]
LIFE_RE = re.compile(r'\bLIFE\b', re.IGNORECASE)


def matches_keywords(text: str) -> bool:
    if not text:
        return False
    tl = text.lower()
    for kw in KEYWORDS:
        if kw.lower() in tl:
            return True
    return bool(LIFE_RE.search(text))


# ─── HTTP session ─────────────────────────────────────────────────────────────
SESSION = requests.Session()
SESSION.headers.update({
    'User-Agent': (
        'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
        'AppleWebKit/537.36 (KHTML, like Gecko) '
        'Chrome/124.0.0.0 Safari/537.36'
    ),
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
    'Accept-Language': 'it-IT,it;q=0.9,en-US;q=0.8',
    'Accept-Encoding': 'gzip, deflate',
    'Connection': 'keep-alive',
})


class GeoBlockedError(Exception):
    pass


def _get(url, params=None, **kw):
    for attempt in range(3):
        try:
            r = SESSION.get(url, params=params, timeout=45, verify=False, **kw)
            if r.status_code == 503:
                raise GeoBlockedError(
                    f'HTTP 503 - Portale non raggiungibile da server cloud '
                    f'(probabile blocco IP non-italiano). '
                    f'Eseguire lo script localmente da PC con connessione italiana.'
                )
            r.raise_for_status()
            return r
        except GeoBlockedError:
            raise
        except Exception:
            if attempt == 2:
                raise
            time.sleep(2 ** attempt)


def _post(url, data=None, **kw):
    for attempt in range(3):
        try:
            r = SESSION.post(url, data=data, timeout=45, verify=False, **kw)
            if r.status_code == 503:
                raise GeoBlockedError(
                    f'HTTP 503 - Portale non raggiungibile da server cloud '
                    f'(probabile blocco IP non-italiano). '
                    f'Eseguire lo script localmente da PC con connessione italiana.'
                )
            r.raise_for_status()
            return r
        except GeoBlockedError:
            raise
        except Exception:
            if attempt == 2:
                raise
            time.sleep(2 ** attempt)


# ─── ASP Albo Online ──────────────────────────────────────────────────────────
ASP_PORTALS = [
    ('ASP Cosenza',        'https://online-aspco.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
    ('ASP Catanzaro',      'https://online-aspcz.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
    ('ASP Crotone',        'https://online-aspkr.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
    ('ASP Reggio Calabria','https://online-asprc.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
    ('ASP Vibo Valentia',  'https://online-aspvv.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
]


def _asp_get_form(soup, base_url):
    """Return (action_url, hidden_fields_dict) from first form found."""
    form = soup.find('form')
    hidden = {}
    action = base_url
    if form:
        action = form.get('action', base_url)
        if not action.startswith('http'):
            action = urljoin(base_url, action)
        for inp in form.find_all('input', type='hidden'):
            n = inp.get('name')
            if n:
                hidden[n] = inp.get('value', '')
    return action, hidden


def _asp_parse_table(soup, base_url):
    """Parse all data rows from the ASP results table.

    Returns list of dicts with keys: data, oggetto, link.
    Also returns total row count (0 = stop pagination).
    """
    table = None
    for attr in [
        {'id': re.compile(r'risultati|tabella|albo|result', re.I)},
        {'class': re.compile(r'risultati|tabella|albo|result|dati', re.I)},
    ]:
        table = soup.find('table', attr)
        if table:
            break
    if not table:
        tables = soup.find_all('table')
        table = tables[0] if tables else None
    if not table:
        return [], 0

    all_rows = table.find_all('tr')
    # Identify header row
    header_row = table.find('tr')
    headers = []
    if header_row:
        headers = [
            c.get_text(strip=True).lower()
            for c in header_row.find_all(['th', 'td'])
        ]

    # Determine column indices from headers
    def col(keywords, default):
        for i, h in enumerate(headers):
            if any(k in h for k in keywords):
                return i
        return default

    date_idx = col(['data pubbl', 'data di pubbl', 'data pubblica', 'data'], 1)
    obj_idx  = col(['oggetto', 'titolo', 'descrizione'], 2)

    data_rows = [r for r in all_rows if r.find('td')]
    results = []
    for row in data_rows:
        cells = row.find_all(['td', 'th'])
        texts = [c.get_text(strip=True) for c in cells]
        if not texts:
            continue

        data_val = texts[date_idx] if date_idx < len(texts) else ''
        oggetto  = texts[obj_idx]  if obj_idx  < len(texts) else ''

        # Find link
        link_url = ''
        for cell in cells:
            a = cell.find('a', href=True)
            if a:
                href = a['href']
                if href and href not in ('#', 'javascript:void(0)', 'javascript:;'):
                    link_url = href if href.startswith('http') else urljoin(base_url, href)
                    break

        results.append({'data': data_val, 'oggetto': oggetto, 'link': link_url})

    return results, len(data_rows)


def _asp_has_next(soup, current_page):
    # Look for "next" / "successivo" link/button that is not disabled
    next_a = soup.find('a', string=re.compile(r'success|avanti|next|»|>', re.I))
    if next_a:
        cl = ' '.join(next_a.get('class', []))
        if 'disabled' not in cl.lower():
            return True
    # Numbered pagination
    for a in soup.find_all('a', string=re.compile(r'^\d+$')):
        try:
            if int(a.get_text(strip=True)) > current_page:
                return True
        except ValueError:
            pass
    return False


def scrape_asp(name, base_url):
    results = []
    error_msg = None
    print(f'\n[{name}] {base_url}')
    try:
        # Initial GET (session init)
        resp = _get(base_url)
        soup = BeautifulSoup(resp.text, 'lxml')
        action, hidden = _asp_get_form(soup, base_url)
        print(f'  form action={action}  hidden={list(hidden.keys())}')

        search_data = {
            **hidden,
            'dataPubblicazioneDal': DATE_FROM_2,
            'dataPubblicazioneAl':  DATE_TO,
        }

        page = 1
        total_scraped = 0

        while True:
            print(f'  Page {page}...', end=' ')
            if page > 1:
                search_data.update({'page': str(page), 'currentPage': str(page)})

            try:
                resp = _post(action, data=search_data)
            except Exception:
                resp = _get(base_url, params={
                    'dataPubblicazioneDal': DATE_FROM_2,
                    'dataPubblicazioneAl': DATE_TO,
                })

            soup = BeautifulSoup(resp.text, 'lxml')
            rows, count = _asp_parse_table(soup, base_url)
            total_scraped += count
            print(f'{count} rows, {sum(1 for r in rows if matches_keywords(r["oggetto"]))} matched')

            for row in rows:
                if matches_keywords(row['oggetto']):
                    results.append(row)

            if count == 0 or not _asp_has_next(soup, page):
                break
            page += 1
            if page > 200:
                break
            time.sleep(0.4)

        print(f'  => {len(results)} matched / {total_scraped} total')

    except GeoBlockedError as e:
        error_msg = str(e)
        print(f'  {error_msg}')
    except Exception as e:
        error_msg = f'ERRORE: {type(e).__name__}: {e}'
        print(f'  {error_msg}')

    return results, error_msg


# ─── Regione Calabria ─────────────────────────────────────────────────────────
def scrape_regione_calabria():
    base_url = 'https://www.regione.calabria.it/provvedimenti-della-regione/'
    results = []
    error_msg = None
    print(f'\n[Regione Calabria] {base_url}')

    try:
        page = 1
        total = 0
        while True:
            params = {
                'filter_date_from': DATE_FROM_2,
                'filter_date_to':   DATE_TO,
            }
            if page > 1:
                params['paged'] = str(page)

            print(f'  Page {page}...', end=' ')
            resp = _get(base_url, params=params)
            soup = BeautifulSoup(resp.text, 'lxml')

            table = soup.find('table')
            if not table:
                print('no table')
                break

            data_rows = [r for r in table.find_all('tr') if r.find('td')]
            if not data_rows:
                break

            print(f'{len(data_rows)} rows', end=', ')

            for row in data_rows:
                # Columns: Tipologia | Data | N(th) | Dipartimento | Oggetto | Dettaglio(link)
                all_cells = row.find_all(['td', 'th'])
                texts = [c.get_text(strip=True) for c in all_cells]

                tipologia  = texts[0] if len(texts) > 0 else ''
                data_val   = texts[1] if len(texts) > 1 else ''
                numero     = texts[2] if len(texts) > 2 else ''
                dipart     = texts[3] if len(texts) > 3 else ''
                oggetto    = texts[4] if len(texts) > 4 else ''

                # Link is in the last cell (Dettaglio)
                link_url = ''
                last_cell = all_cells[-1] if all_cells else None
                if last_cell:
                    a = last_cell.find('a', href=True)
                    if a:
                        link_url = a['href']
                        if not link_url.startswith('http'):
                            link_url = urljoin(base_url, link_url)

                total += 1
                if matches_keywords(oggetto):
                    results.append({
                        'data':    data_val,
                        'oggetto': oggetto,
                        'link':    link_url,
                    })

            matched_page = sum(1 for r in results)
            print(f'{len(results)} matched so far')

            # Pagination
            next_a = soup.find('a', class_='next page-numbers')
            if not next_a:
                break
            page += 1
            if page > 500:
                break
            time.sleep(0.3)

        print(f'  => {len(results)} matched / {total} total')

    except GeoBlockedError as e:
        error_msg = str(e)
        print(f'  {error_msg}')
    except Exception as e:
        error_msg = f'ERRORE: {type(e).__name__}: {e}'
        print(f'  {error_msg}')

    return results, error_msg


# ─── TAR Catanzaro ───────────────────────────────────────────────────────────
def scrape_tar_catanzaro():
    """TAR Catanzaro - filter on 'Parte' column with publishDateFrom last 3 days."""
    base_url = 'https://www.giustizia-amministrativa.it/provvedimenti-tar-catanzaro'
    results = []
    error_msg = None
    print(f'\n[TAR Catanzaro] {base_url}')

    try:
        page = 0
        total = 0
        while True:
            params = {
                'publishDateFrom': DATE_FROM_3,
                'publishDateTo':   DATE_TO,
            }
            if page > 0:
                params['cur']   = str(page + 1)
                params['delta'] = '20'

            print(f'  Page {page + 1}...', end=' ')
            resp = _get(base_url, params=params)
            soup = BeautifulSoup(resp.text, 'lxml')

            table = soup.find('table')
            if not table:
                print('no table')
                break

            # Identify headers
            header_row = table.find('tr')
            headers = []
            if header_row:
                headers = [
                    c.get_text(strip=True).lower()
                    for c in header_row.find_all(['th', 'td'])
                ]

            def col(kws, default):
                for i, h in enumerate(headers):
                    if any(k in h for k in kws):
                        return i
                return default

            date_idx  = col(['data'],  0)
            parte_idx = col(['parte'], 3)

            data_rows = [r for r in table.find_all('tr') if r.find('td')]
            if not data_rows:
                break

            print(f'{len(data_rows)} rows (parte_col={parte_idx})', end=', ')

            for row in data_rows:
                cells = row.find_all(['td', 'th'])
                texts = [c.get_text(strip=True) for c in cells]
                if not texts:
                    continue

                data_val = texts[date_idx]  if date_idx  < len(texts) else ''
                parte    = texts[parte_idx] if parte_idx < len(texts) else ''

                link_url = ''
                for cell in cells:
                    a = cell.find('a', href=True)
                    if a and a['href'] not in ('#', ''):
                        href = a['href']
                        link_url = href if href.startswith('http') else urljoin(base_url, href)
                        break

                total += 1
                if matches_keywords(parte):
                    results.append({'data': data_val, 'oggetto': parte, 'link': link_url})

            print(f'{len(results)} matched so far')

            if not _asp_has_next(soup, page + 1):
                break
            page += 1
            if page > 200:
                break
            time.sleep(0.3)

        print(f'  => {len(results)} matched / {total} total')

    except GeoBlockedError as e:
        error_msg = str(e)
        print(f'  {error_msg}')
    except Exception as e:
        error_msg = f'ERRORE: {type(e).__name__}: {e}'
        print(f'  {error_msg}')

    return results, error_msg


# ─── Excel builder ────────────────────────────────────────────────────────────
HDR_FONT  = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
HDR_FILL  = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
HDR_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LINK_FONT = Font(color='0563C1', underline='single', name='Calibri', size=11)
ERR_FONT  = Font(color='C00000', italic=True, name='Calibri', size=11)
WRAP      = Alignment(wrap_text=True, vertical='top')


def _write_sheet(ws, rows, error_msg):
    if error_msg:
        ws.column_dimensions['A'].width = 90
        ws.row_dimensions[1].height = 60
        cell = ws['A1']
        cell.value     = error_msg
        cell.font      = ERR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        return

    if not rows:
        ws['A1'].value = 'Nessun risultato nel periodo'
        return

    headers = ['Data', 'Oggetto', 'Descrizione breve (150 car.)', 'Link']
    col_widths = [16, 70, 45, 12]

    for ci, (hdr, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=ci, value=hdr)
        cell.font      = HDR_FONT
        cell.fill      = HDR_FILL
        cell.alignment = HDR_ALIGN
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.row_dimensions[1].height = 22

    for ri, item in enumerate(rows, 2):
        data_val = item.get('data', '')
        oggetto  = item.get('oggetto', '')
        descr    = oggetto[:150] if oggetto else ''
        link_url = item.get('link', '')

        ws.cell(row=ri, column=1, value=data_val).alignment = WRAP
        obj_cell = ws.cell(row=ri, column=2, value=oggetto)
        obj_cell.alignment = WRAP
        ws.cell(row=ri, column=3, value=descr).alignment = WRAP

        link_cell = ws.cell(row=ri, column=4)
        if link_url:
            link_cell.value     = 'Apri'
            link_cell.hyperlink = link_url
            link_cell.font      = LINK_FONT
            link_cell.alignment = Alignment(horizontal='center', vertical='top')
        else:
            link_cell.value = ''

    # Auto row height hint (approx)
    for ri in range(2, len(rows) + 2):
        ws.row_dimensions[ri].height = 40


def build_excel(all_data, output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # delete default sheet

    for sheet_name, rows, error_msg in all_data:
        safe_name = sheet_name[:31]
        ws = wb.create_sheet(title=safe_name)
        _write_sheet(ws, rows, error_msg)
        status = f'{len(rows)} risultati' if not error_msg else 'ERRORE'
        print(f'  Sheet "{safe_name}": {status}')

    wb.save(output_path)
    print(f'\n✓ Salvato: {output_path}')


# ─── Main ─────────────────────────────────────────────────────────────────────
def main():
    print('=' * 60)
    print('MONITORAGGIO ATTI CALABRIA')
    print(f'Data esecuzione : {TODAY.strftime("%d/%m/%Y %H:%M")}')
    print(f'Periodo ASP/RC  : {DATE_FROM_2} → {DATE_TO}')
    print(f'Periodo TAR     : {DATE_FROM_3} → {DATE_TO}')
    print('=' * 60)

    all_data = []  # list of (sheet_name, rows, error_msg)

    # ASP portals
    for name, url in ASP_PORTALS:
        rows, err = scrape_asp(name, url)
        all_data.append((name, rows, err))

    # Regione Calabria
    rows, err = scrape_regione_calabria()
    all_data.append(('Regione Calabria', rows, err))

    # TAR Catanzaro
    rows, err = scrape_tar_catanzaro()
    all_data.append(('TAR Catanzaro', rows, err))

    # Build Excel
    # OUTPUT PATH - modifica se necessario
    output_path = rf'C:\Users\aoedo\kDrive\01_Lavoro\Clienti\Starbene\Atti\Monitoraggio_Atti_Calabria_{FILE_DATE}.xlsx'
    print('\n=== Generazione Excel ===')
    build_excel(all_data, output_path)

    print('\n=== RIEPILOGO ===')
    for name, rows, err in all_data:
        if err:
            print(f'  {name:<25} ERRORE ({err[:60]})')
        else:
            print(f'  {name:<25} {len(rows)} risultati con keyword match')
    print(f'\nFile: {output_path}')


if __name__ == '__main__':
    main()
