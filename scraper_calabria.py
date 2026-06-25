#!/usr/bin/env python3
"""
Monitoraggio Atti Calabria
Scraper multi-portale: ASP (5), Regione Calabria, TAR Catanzaro
"""

import asyncio
import re
import os
import sys
from datetime import datetime, timedelta

from playwright.async_api import async_playwright, TimeoutError as PWTimeout
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── DATE ─────────────────────────────────────────────────────────────────────
TODAY = datetime(2026, 6, 25)
DATE_FROM_2D = (TODAY - timedelta(days=2)).strftime("%Y-%m-%d")   # 2026-06-23
DATE_FROM_3D = (TODAY - timedelta(days=3)).strftime("%Y-%m-%d")   # 2026-06-22

OUTPUT_DIR  = "/home/user/NS"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, f"Monitoraggio_Atti_Calabria_{TODAY.strftime('%d-%m-%Y')}.xlsx")

# ── KEYWORD MATCHING ──────────────────────────────────────────────────────────
KEYWORDS_PLAIN = [
    "ADI", "Assistenza domiciliare", "ANMIC", "Accreditamento",
    "Aumento di budget", "Autismo", "Autorizzazione all'esercizio",
    "Autorizzazione alla realizzazione", "Autorizzazioni", "Budget",
    "Casa Giardino", "Centro San Giuseppe", "Centro salute e benessere",
    "Fabbisogni LEA", "Fisiolab", "Fisioterapia",
    "Parere commissione", "Presa d'atto verifica", "Programmazione",
    "Rete riabilitativa", "Rete territoriale", "Riabilitazione estensiva",
    "Riconversione prestazioni", "Rinnovo accreditamento",
    "San Teodoro", "Savelli Hospital", "Starbene",
    "Verifica requisiti", "Villa San Giuseppe", "Villa del Rosario",
]
LIFE_RE = re.compile(r'\bLIFE\b', re.IGNORECASE)

def keyword_match(text: str) -> bool:
    if not text:
        return False
    tu = text.upper()
    for kw in KEYWORDS_PLAIN:
        if kw.upper() in tu:
            return True
    return bool(LIFE_RE.search(text))

# ── HELPERS ───────────────────────────────────────────────────────────────────
def make_abs_link(href: str, base_url: str) -> str:
    if not href:
        return ""
    if href.startswith("http"):
        return href
    from urllib.parse import urljoin
    return urljoin(base_url, href)

async def safe_text(locator) -> str:
    try:
        return (await locator.text_content(timeout=5000) or "").strip()
    except Exception:
        return ""

async def safe_attr(locator, attr: str) -> str:
    try:
        return (await locator.get_attribute(attr, timeout=5000) or "").strip()
    except Exception:
        return ""

# ── ASP PORTALI ──────────────────────────────────────────────────────────────
ASP_PORTALS = [
    ("ASP Cosenza",        "https://online-aspco.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
    ("ASP Catanzaro",      "https://online-aspcz.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
    ("ASP Crotone",        "https://online-aspkr.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
    ("ASP Reggio Calabria","https://online-asprc.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
    ("ASP Vibo Valentia",  "https://online-aspvv.sisr.regione.calabria.it/AlboOnline/ricercaAlbo"),
]

async def scrape_asp(browser, name: str, url: str) -> list:
    results = []
    page = await browser.new_page()
    try:
        print(f"\n[{name}] → {url}")
        await page.goto(url, wait_until="domcontentloaded", timeout=45000)
        await page.wait_for_timeout(2000)

        # fill date "dal"
        try:
            inp = page.locator("#dataPubblicazioneDal")
            if await inp.count():
                await inp.fill(DATE_FROM_2D)
                print(f"[{name}] Set dataPubblicazioneDal = {DATE_FROM_2D}")
        except Exception as e:
            print(f"[{name}] Date field error: {e}")

        # submit
        for sel in ['button[type="submit"]', 'input[type="submit"]', 'a:has-text("Cerca")', 'button:has-text("Cerca")']:
            btn = page.locator(sel).first
            if await btn.count():
                await btn.click()
                break
        await page.wait_for_load_state("networkidle", timeout=30000)
        await page.wait_for_timeout(1500)

        page_num = 0
        while True:
            page_num += 1
            print(f"[{name}] Page {page_num} – extracting rows…")

            # Try multiple table selectors
            rows = []
            for tbl_sel in ["table tbody tr", "#risultati tbody tr", ".albo-table tbody tr", "table.table tbody tr"]:
                rows = await page.locator(tbl_sel).all()
                if rows:
                    break

            if not rows:
                # Try any visible rows
                rows = await page.locator("tr").all()

            extracted = 0
            for row in rows:
                cells = await row.locator("td").all()
                if len(cells) < 2:
                    continue

                # Column heuristic: date is usually first numeric column
                date_text = await safe_text(cells[0])
                # Try to find the oggetto column (often column 2 or 3)
                oggetto = ""
                link_href = ""
                for i, cell in enumerate(cells):
                    t = await safe_text(cell)
                    # oggetto tends to be the longest text
                    if len(t) > len(oggetto):
                        oggetto = t
                    # grab first link found
                    if not link_href:
                        a = cell.locator("a").first
                        if await a.count():
                            link_href = await safe_attr(a, "href")

                link_href = make_abs_link(link_href, url)
                if keyword_match(oggetto):
                    results.append({"data": date_text, "oggetto": oggetto, "link": link_href})
                    extracted += 1

            print(f"[{name}] Page {page_num}: {extracted} matches so far (total {len(results)})")

            # Navigate to next page
            next_found = False
            for next_sel in [
                "a:has-text('Successivo')", "a:has-text('Avanti')",
                "a[aria-label='Pagina successiva']", ".pagNext a", ".next a",
                "a.pagNext", "li.next a", "a:has-text('>')",
            ]:
                nxt = page.locator(next_sel).first
                if await nxt.count() and await nxt.is_visible():
                    try:
                        await nxt.click()
                        await page.wait_for_load_state("networkidle", timeout=20000)
                        await page.wait_for_timeout(1000)
                        next_found = True
                        break
                    except Exception:
                        pass
            if not next_found:
                break

    except Exception as e:
        print(f"[{name}] FATAL: {e}")
    finally:
        await page.close()

    print(f"[{name}] Done – {len(results)} matching records")
    return results


# ── REGIONE CALABRIA ──────────────────────────────────────────────────────────
RC_URL = "https://www.regione.calabria.it/provvedimenti-della-regione/"

async def scrape_regione(browser) -> list:
    results = []
    page = await browser.new_page()
    name = "Regione Calabria"
    try:
        print(f"\n[{name}] → {RC_URL}")
        await page.goto(RC_URL, wait_until="domcontentloaded", timeout=45000)
        await page.wait_for_timeout(2000)

        # Try date filter
        for date_sel in ["#filter_date_from", "input[name='filter_date_from']",
                          "input[placeholder*='dal']", "input[type='date']"]:
            inp = page.locator(date_sel).first
            if await inp.count():
                await inp.fill(DATE_FROM_2D)
                print(f"[{name}] Set date from = {DATE_FROM_2D}")
                break

        # Submit/filter
        for btn_sel in ['button[type="submit"]', 'input[type="submit"]',
                         'button:has-text("Filtra")', 'button:has-text("Cerca")']:
            btn = page.locator(btn_sel).first
            if await btn.count():
                await btn.click()
                break
        await page.wait_for_load_state("networkidle", timeout=30000)
        await page.wait_for_timeout(1500)

        page_num = 0
        while True:
            page_num += 1
            print(f"[{name}] Page {page_num}…")

            # Try table rows
            rows = []
            for sel in ["table tbody tr", ".view-content .views-row",
                         "article", ".provvedimento-row", "tr"]:
                rows = await page.locator(sel).all()
                if rows:
                    break

            extracted = 0
            for row in rows:
                text = await safe_text(row)
                if not text or len(text) < 5:
                    continue
                # Extract date (look for cell or span with date)
                date_text = ""
                for date_lbl in [".date", ".field-date", "td:first-child", "time"]:
                    dt = row.locator(date_lbl).first
                    if await dt.count():
                        date_text = await safe_text(dt)
                        break

                # Extract oggetto/title
                oggetto = ""
                for obj_sel in [".title", ".oggetto", "td:nth-child(2)", "h3", "h2", ".views-field-title"]:
                    el = row.locator(obj_sel).first
                    if await el.count():
                        t = await safe_text(el)
                        if t:
                            oggetto = t
                            break
                if not oggetto:
                    oggetto = text[:300]

                # Link
                link_href = ""
                a = row.locator("a").first
                if await a.count():
                    link_href = await safe_attr(a, "href")
                link_href = make_abs_link(link_href, RC_URL)

                if keyword_match(oggetto):
                    results.append({"data": date_text, "oggetto": oggetto, "link": link_href})
                    extracted += 1

            print(f"[{name}] Page {page_num}: {extracted} matches (total {len(results)})")

            # Next page
            nxt = page.locator("a:has-text('Successivo'), .pager-next a, li.next a, a.next, a[rel='next']").first
            if await nxt.count() and await nxt.is_visible():
                try:
                    await nxt.click()
                    await page.wait_for_load_state("networkidle", timeout=20000)
                    await page.wait_for_timeout(1000)
                except Exception:
                    break
            else:
                break

    except Exception as e:
        print(f"[{name}] FATAL: {e}")
    finally:
        await page.close()

    print(f"[{name}] Done – {len(results)} matching records")
    return results


# ── TAR CATANZARO ─────────────────────────────────────────────────────────────
TAR_URL = "https://www.giustizia-amministrativa.it/provvedimenti-tar-catanzaro"

async def scrape_tar(browser) -> list:
    results = []
    page = await browser.new_page()
    name = "TAR Catanzaro"
    try:
        print(f"\n[{name}] → {TAR_URL}")
        await page.goto(TAR_URL, wait_until="domcontentloaded", timeout=45000)
        await page.wait_for_timeout(3000)

        # Set publishDateFrom
        for sel in ["#publishDateFrom", "input[name='publishDateFrom']",
                     "input[placeholder*='data']", "input[type='date']"]:
            inp = page.locator(sel).first
            if await inp.count():
                await inp.fill(DATE_FROM_3D)
                print(f"[{name}] Set publishDateFrom = {DATE_FROM_3D}")
                break

        # Submit
        for btn_sel in ['button[type="submit"]', 'input[type="submit"]',
                          'button:has-text("Cerca")', 'button:has-text("Filtra")']:
            btn = page.locator(btn_sel).first
            if await btn.count():
                await btn.click()
                break
        await page.wait_for_load_state("networkidle", timeout=30000)
        await page.wait_for_timeout(2000)

        page_num = 0
        while True:
            page_num += 1
            print(f"[{name}] Page {page_num}…")

            rows = []
            for sel in ["table tbody tr", ".provvedimento", ".row-item", "tr"]:
                rows = await page.locator(sel).all()
                if rows:
                    break

            extracted = 0
            for row in rows:
                cells = await row.locator("td").all()
                text = await safe_text(row)
                if not text or len(text) < 3:
                    continue

                date_text = ""
                parte = ""
                link_href = ""

                if cells:
                    date_text = await safe_text(cells[0])
                    # "Parte" column for TAR - check all cells
                    for cell in cells:
                        t = await safe_text(cell)
                        if t and t != date_text:
                            parte = t
                            break

                # link
                a = row.locator("a").first
                if await a.count():
                    link_href = await safe_attr(a, "href")
                link_href = make_abs_link(link_href, TAR_URL)

                # For TAR filter on "Parte" column (many are XXX_OMISSIS_XXX)
                if keyword_match(parte) or keyword_match(text):
                    results.append({"data": date_text, "oggetto": parte or text[:300], "link": link_href})
                    extracted += 1

            print(f"[{name}] Page {page_num}: {extracted} matches (total {len(results)})")

            nxt = page.locator("a:has-text('Successivo'), .pager-next a, li.next a, a[rel='next']").first
            if await nxt.count() and await nxt.is_visible():
                try:
                    await nxt.click()
                    await page.wait_for_load_state("networkidle", timeout=20000)
                    await page.wait_for_timeout(1000)
                except Exception:
                    break
            else:
                break

    except Exception as e:
        print(f"[{name}] FATAL: {e}")
    finally:
        await page.close()

    print(f"[{name}] Done – {len(results)} matching records")
    return results


# ── XLSX EXPORT ───────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
LINK_FONT   = Font(color="0563C1", underline="single", name="Calibri", size=10)
CELL_FONT   = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

def write_xlsx(all_data: dict):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    sheets_order = [
        "ASP Cosenza", "ASP Catanzaro", "ASP Crotone",
        "ASP Reggio Calabria", "ASP Vibo Valentia",
        "Regione Calabria", "TAR Catanzaro",
    ]

    for sheet_name in sheets_order:
        rows = all_data.get(sheet_name, [])
        ws = wb.create_sheet(title=sheet_name[:31])

        if not rows:
            ws["A1"] = "Nessun risultato"
            ws["A1"].font = Font(italic=True, color="808080")
            continue

        # Header
        headers = ["Data", "Oggetto", "Descrizione breve (150 car.)", "Link"]
        col_widths = [18, 60, 55, 50]
        for col_idx, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER
            ws.column_dimensions[get_column_letter(col_idx)].width = w
        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "A2"

        # Data rows
        for r_idx, rec in enumerate(rows, 2):
            data    = rec.get("data", "")
            oggetto = rec.get("oggetto", "")
            desc    = oggetto[:150] if oggetto else ""
            link    = rec.get("link", "")

            cells_vals = [data, oggetto, desc, link]
            for c_idx, val in enumerate(cells_vals, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = CELL_FONT
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=(c_idx == 2))

            # Hyperlink in col 4
            if link:
                link_cell = ws.cell(row=r_idx, column=4)
                link_cell.value = link
                link_cell.hyperlink = link
                link_cell.font = LINK_FONT

        ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_FILE)
    print(f"\n✅ File saved: {OUTPUT_FILE}")


# ── MAIN ─────────────────────────────────────────────────────────────────────
async def main():
    print(f"Date range 2-day: {DATE_FROM_2D} → {TODAY.strftime('%Y-%m-%d')}")
    print(f"Date range 3-day: {DATE_FROM_3D} → {TODAY.strftime('%Y-%m-%d')}")
    print(f"Output: {OUTPUT_FILE}\n")

    all_data = {}

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(
            executable_path="/opt/pw-browsers/chromium",
            headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage", "--disable-blink-features=AutomationControlled"],
        )

        # ASP portals (sequential to avoid hammering)
        for portal_name, portal_url in ASP_PORTALS:
            res = await scrape_asp(browser, portal_name, portal_url)
            all_data[portal_name] = res

        # Regione Calabria
        res = await scrape_regione(browser)
        all_data["Regione Calabria"] = res

        # TAR Catanzaro
        res = await scrape_tar(browser)
        all_data["TAR Catanzaro"] = res

        await browser.close()

    # Summary
    print("\n=== SUMMARY ===")
    for sheet, rows in all_data.items():
        print(f"  {sheet}: {len(rows)} matches")

    write_xlsx(all_data)


if __name__ == "__main__":
    asyncio.run(main())
