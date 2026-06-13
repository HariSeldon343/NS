#!/usr/bin/env python3
"""
Monitoraggio Atti Calabria - v2
Portali: 5 ASP, Regione Calabria, TAR Catanzaro
"""

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import re, json, time, urllib3
from datetime import datetime, timedelta
from urllib.parse import urljoin, urlparse

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ── Dates ─────────────────────────────────────────────────────────────────────
TODAY       = datetime.now()
DATE_FROM_2 = (TODAY - timedelta(days=2)).strftime('%Y-%m-%d')
DATE_FROM_3 = (TODAY - timedelta(days=3)).strftime('%Y-%m-%d')
DATE_TO     = TODAY.strftime('%Y-%m-%d')
FILE_DATE   = TODAY.strftime('%d-%m-%Y')
OUTPUT_FILE = f'/home/user/NS/Monitoraggio_Atti_Calabria_{FILE_DATE}.xlsx'

print(f"Data esecuzione : {TODAY.strftime('%d/%m/%Y %H:%M')}")
print(f"Intervallo 2gg  : {DATE_FROM_2} → {DATE_TO}")
print(f"Intervallo 3gg  : {DATE_FROM_3} → {DATE_TO}")
print(f"Output          : {OUTPUT_FILE}\n")

# ── Keywords ─────────────────────────────────────────────────────────────────
KEYWORDS = [
    'ADI', 'Assistenza domiciliare', 'ANMIC', 'Accreditamento',
    'Aumento di budget', 'Autismo', "Autorizzazione all'esercizio",
    'Autorizzazione alla realizzazione', 'Autorizzazioni', 'Budget',
    'Casa Giardino', 'Centro San Giuseppe', 'Centro salute e benessere',
    'Fabbisogni LEA', 'Fisiolab', 'Fisioterapia', 'Parere commissione',
    "Presa d'atto verifica", 'Programmazione', 'Rete riabilitativa',
    'Rete territoriale', 'Riabilitazione estensiva', 'Riconversione prestazioni',
    'Rinnovo accreditamento', 'San Teodoro', 'Savelli Hospital', 'Starbene',
    'Verifica requisiti', 'Villa San Giuseppe', 'Villa del Rosario',
]
LIFE_RE = re.compile(r'\bLIFE\b', re.IGNORECASE)

def matches_keywords(text):
    if not text:
        return False
    tu = text.upper()
    for kw in KEYWORDS:
        if kw.upper() in tu:
            return True
    return bool(LIFE_RE.search(text))

# ── HTTP helpers ──────────────────────────────────────────────────────────────
CHROME_HEADERS = {
    'User-Agent': (
        'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
        'AppleWebKit/537.36 (KHTML, like Gecko) '
        'Chrome/124.0.0.0 Safari/537.36'
    ),
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
    'Accept-Language': 'it-IT,it;q=0.9,en-US;q=0.5',
}

def make_session(extra_headers=None):
    s = requests.Session()
    s.headers.update(CHROME_HEADERS)
    if extra_headers:
        s.headers.update(extra_headers)
    return s

# ── ASP portals ───────────────────────────────────────────────────────────────
ASP_PORTALS = [
    ('ASP Cosenza',         'https://online-aspco.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
    ('ASP Catanzaro',       'https://online-aspcz.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
    ('ASP Crotone',         'https://online-aspkr.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
    ('ASP Reggio Calabria', 'https://online-asprc.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
    ('ASP Vibo Valentia',   'https://online-aspvv.sisr.regione.calabria.it/AlboOnline/ricercaAlbo'),
]

def _collect_form(form, url):
    action = form.get('action') or url
    if not action.startswith('http'):
        action = urljoin(url, action)
    method = (form.get('method') or 'post').lower()
    data = {}
    for tag in form.find_all(['input', 'select', 'textarea']):
        name = tag.get('name') or ''
        val  = tag.get('value', '')
        typ  = tag.get('type', 'text').lower()
        if name and typ not in ('submit', 'button', 'image', 'reset'):
            if tag.name == 'select':
                opt = tag.find('option', selected=True)
                val = opt.get('value', '') if opt else ''
            data[name] = val
    return action, method, data

def _parse_table(soup, base_url):
    """Return list of {data, oggetto, link, _row} from best matching table."""
    rows = []
    best = None
    for tbl in soup.find_all('table'):
        ths = [th.get_text(strip=True).lower() for th in tbl.find_all('th')]
        if any(h in ths for h in ('oggetto', 'data', 'numero', 'tipo', 'atto', 'descrizione')):
            best = tbl
            break
        body = tbl.get_text().lower()
        if 'oggetto' in body or ('data' in body and 'numero' in body):
            best = tbl
    if best is None:
        tables = soup.find_all('table')
        if tables:
            best = tables[0]
    if best is None:
        return rows

    all_tr = best.find_all('tr')
    if not all_tr:
        return rows
    hdrs = [c.get_text(strip=True).lower() for c in all_tr[0].find_all(['th', 'td'])]

    def ci(pats):
        for p in pats:
            for i, h in enumerate(hdrs):
                if p in h:
                    return i
        return -1

    dc = ci(['data', 'date'])
    oc = ci(['oggetto', 'descrizione', 'titolo', 'object'])

    for tr in all_tr[1:]:
        cells = tr.find_all(['td', 'th'])
        if not cells:
            continue
        texts = [c.get_text(strip=True) for c in cells]
        data_val    = texts[dc] if 0 <= dc < len(texts) else (texts[0] if texts else '')
        oggetto_val = texts[oc] if 0 <= oc < len(texts) else ' | '.join(texts)
        link = None
        for c in cells:
            a = c.find('a', href=True)
            if a and a['href'] and not a['href'].startswith('#'):
                link = urljoin(base_url, a['href'])
                break
        rows.append({'data': data_val, 'oggetto': oggetto_val, 'link': link or base_url,
                     '_row': ' | '.join(texts)})
    return rows

def scrape_asp(name, base_url):
    results = []
    error_msg = None
    s = make_session()
    print(f"\n{'='*60}\n[{name}] {base_url}")

    try:
        r = s.get(base_url, timeout=25, verify=False)
        print(f"[{name}] GET → {r.status_code} ({len(r.content)} B)")

        if r.status_code != 200:
            error_msg = f"Portale non raggiungibile – HTTP {r.status_code}"
            print(f"[{name}] {error_msg}")
            return results, error_msg

        soup = BeautifulSoup(r.text, 'lxml')
        form = soup.find('form')
        if not form:
            # Try direct parse (no form = already a listing)
            rows = _parse_table(soup, base_url)
            for row in rows:
                if matches_keywords(row['oggetto']) or matches_keywords(row['_row']):
                    results.append({'data': row['data'], 'oggetto': row['oggetto'], 'link': row['link']})
            print(f"[{name}] No form – direct parse: {len(results)} matches")
            return results, None

        form_action, form_method, form_data = _collect_form(form, base_url)
        print(f"[{name}] Form: {form_method.upper()} {form_action[:70]}")
        print(f"[{name}] Fields: {list(form_data.keys())[:10]}")

        # Inject date
        set_from = False
        for key in list(form_data.keys()):
            lk = key.lower()
            if 'datapubblicazionedal' in lk or lk in ('dal', 'datadal', 'datafrom'):
                form_data[key] = DATE_FROM_2
                set_from = True
            elif 'datapubblicazioneal' in lk or lk in ('al', 'dataal', 'datato'):
                form_data[key] = DATE_TO
        if not set_from:
            form_data['dataPubblicazioneDal'] = DATE_FROM_2
            form_data['dataPubblicazioneAl']  = DATE_TO

        page = 1
        while True:
            if form_method == 'post':
                r2 = s.post(form_action, data=form_data, timeout=25, verify=False)
            else:
                r2 = s.get(form_action, params=form_data, timeout=25, verify=False)
            print(f"[{name}] Page {page} → {r2.status_code} ({len(r2.content)} B)")
            if r2.status_code != 200:
                break

            soup2 = BeautifulSoup(r2.text, 'lxml')
            rows  = _parse_table(soup2, base_url)
            print(f"[{name}] Page {page}: {len(rows)} rows in table")
            if not rows:
                break

            for row in rows:
                if matches_keywords(row['oggetto']) or matches_keywords(row['_row']):
                    results.append({'data': row['data'], 'oggetto': row['oggetto'], 'link': row['link']})

            # next page
            next_url = None
            for a in soup2.find_all('a', href=True):
                txt = a.get_text(strip=True)
                if txt in ('Successiva', '>', '»', 'Next', 'Avanti', '›'):
                    next_url = urljoin(base_url, a['href'])
                    break
            nl = soup2.find('a', rel='next')
            if nl and nl.get('href'):
                next_url = urljoin(base_url, nl['href'])
            if not next_url or next_url == form_action:
                break
            form_action = next_url
            form_method = 'get'
            form_data   = {}
            page += 1
            time.sleep(0.4)

    except requests.exceptions.Timeout:
        error_msg = "Portale non raggiungibile – timeout di connessione"
        print(f"[{name}] {error_msg}")
    except requests.exceptions.ConnectionError as e:
        error_msg = f"Portale non raggiungibile – errore di rete"
        print(f"[{name}] {error_msg}: {e}")
    except Exception as e:
        import traceback
        error_msg = f"Errore imprevisto: {e}"
        print(f"[{name}] {error_msg}")
        traceback.print_exc()

    print(f"[{name}] Matches: {len(results)}")
    return results, error_msg

# ── Regione Calabria ──────────────────────────────────────────────────────────

def scrape_regione_calabria():
    results = []
    base_url = 'https://www.regione.calabria.it/provvedimenti-della-regione/'
    s = make_session()
    print(f"\n{'='*60}\n[Regione Calabria] {base_url}")

    try:
        page = 1
        total_rows = 0
        while True:
            params = {'filter_date_from': DATE_FROM_2}
            if page > 1:
                params['paged'] = str(page)

            r = s.get(base_url, params=params, timeout=30, verify=False)
            print(f"[Regione Calabria] Page {page} → {r.status_code} ({len(r.content)} B)")
            if r.status_code != 200:
                break

            soup = BeautifulSoup(r.text, 'lxml')
            tbl  = soup.find('table')
            if not tbl:
                break

            all_tr = tbl.find_all('tr')
            hdrs   = [c.get_text(strip=True).lower() for c in all_tr[0].find_all(['th','td'])] if all_tr else []

            def ci2(pats):
                for p in pats:
                    for i, h in enumerate(hdrs):
                        if p in h:
                            return i
                return -1

            dc = ci2(['data', 'repertoriazione'])
            oc = ci2(['oggetto', 'descrizione', 'titolo'])
            # columns typically: Tipologia | Data | N | Dipartimento | Oggetto | Dettaglio

            page_count = 0
            for tr in all_tr[1:]:
                cells = tr.find_all(['td', 'th'])
                if not cells:
                    continue
                page_count += 1
                total_rows  += 1
                texts = [c.get_text(strip=True) for c in cells]
                data_val    = texts[dc] if 0 <= dc < len(texts) else (texts[1] if len(texts) > 1 else '')
                oggetto_val = texts[oc] if 0 <= oc < len(texts) else (texts[4] if len(texts) > 4 else ' | '.join(texts))
                link = None
                for c in cells:
                    a = c.find('a', href=True)
                    if a and a['href'] and not a['href'].startswith('#'):
                        link = urljoin(base_url, a['href'])
                        break

                if matches_keywords(oggetto_val):
                    print(f"[Regione Calabria] MATCH [{data_val}]: {oggetto_val[:100]}")
                    results.append({'data': data_val, 'oggetto': oggetto_val, 'link': link or base_url})

            print(f"[Regione Calabria] Page {page}: {page_count} rows, total so far {total_rows}")
            if page_count == 0:
                break

            # Pagination
            next_url = None
            nl = soup.find('a', rel='next')
            if nl and nl.get('href'):
                next_url = urljoin(base_url, nl['href'])
            if not next_url:
                for a in soup.find_all('a', href=True):
                    if f'paged={page+1}' in a['href']:
                        next_url = urljoin(base_url, a['href'])
                        break
            if not next_url:
                break
            page += 1
            time.sleep(0.3)

    except Exception as e:
        import traceback
        print(f"[Regione Calabria] ERROR: {e}")
        traceback.print_exc()

    print(f"[Regione Calabria] Matches: {len(results)}")
    return results

# ── TAR Catanzaro (AJAX / DataTables) ────────────────────────────────────────

TAR_PORTLET_INSTANCE = 'jjYpzZYF4Qfe'
TAR_PREFIX = (
    '_it_indra_ga_institutional_area_'
    'JurisdictionalActivityAdministrativeActsWebPortlet_'
    f'INSTANCE_{TAR_PORTLET_INSTANCE}_'
)
TAR_BASE   = 'https://www.giustizia-amministrativa.it/provvedimenti-tar-catanzaro'
TAR_SEARCH = (
    'https://www.giustizia-amministrativa.it/web/guest/provvedimenti-tar-catanzaro'
    '?p_p_id=it_indra_ga_institutional_area_'
    f'JurisdictionalActivityAdministrativeActsWebPortlet_INSTANCE_{TAR_PORTLET_INSTANCE}'
    '&p_p_lifecycle=2&p_p_state=normal&p_p_mode=view'
    '&p_p_resource_id=%2Fadministrative-acts%2Fsearch%2Fresults'
    '&p_p_cacheability=cacheLevelPage'
)
TAR_COLUMNS = [
    'nrgFascicolo','sezione','parte','tipoUdienza','dataUdienza',
    'numProvvedimento','dataPubblicazione','tipoProvvedimento',
    'relatore','presidente','esito',
]

def _tar_get_form_date(s):
    """Load TAR page and extract form_date token."""
    r = s.get(TAR_BASE, timeout=30, verify=False)
    soup = BeautifulSoup(r.text, 'lxml')
    for form in soup.find_all('form'):
        for inp in form.find_all('input'):
            n = inp.get('name', '')
            if 'formDate' in n:
                return inp.get('value', '')
    return str(int(datetime.now().timestamp() * 1000))

def _tar_post_form(s, form_date):
    """Submit search form to get session/cookie for AJAX."""
    # Get form action
    r0 = s.get(TAR_BASE, timeout=30, verify=False)
    soup0 = BeautifulSoup(r0.text, 'lxml')
    form_action = TAR_BASE
    for form in soup0.find_all('form'):
        if any('publishDate' in (i.get('name','')) for i in form.find_all('input')):
            form_action = form.get('action', TAR_BASE)
            break
    post_data = {
        TAR_PREFIX + 'formDate':          form_date,
        TAR_PREFIX + 'publishDateFrom':   DATE_FROM_3,
        TAR_PREFIX + 'publishDateTo':     DATE_TO,
        TAR_PREFIX + 'number':            '',
        TAR_PREFIX + 'president':         '',
        TAR_PREFIX + 'draftingJudge':     '',
        TAR_PREFIX + 'year':              '',
        TAR_PREFIX + 'section':           '',
        TAR_PREFIX + 'type':              '',
        TAR_PREFIX + 'specific':          '',
    }
    s.post(form_action, data=post_data, timeout=30, verify=False)

def _tar_ajax_search(s, start=0, length=200):
    """Call TAR DataTables AJAX endpoint and return JSON."""
    additional_info = json.dumps({
        'schema': 'TAR_CATANZARO', 'type': None, 'year': '', 'number': '',
        'hearingDateFrom': None, 'hearingDateTo': None,
        'publishDateFrom': DATE_FROM_3, 'publishDateTo': DATE_TO,
        'hearingType': None, 'nrg': None, 'section': '',
        'provisionSpecification': '', 'president': '', 'draftingJudge': '',
        'subjectMatter': None, 'page': None, 'size': None,
        'orderBy': None, 'orderStrategy': None, 'queryString': None,
    })
    payload = json.dumps({
        'draw': 1,
        'columns': [
            {'data': c, 'name': '', 'searchable': True, 'orderable': True,
             'search': {'value': '', 'regex': False}}
            for c in TAR_COLUMNS
        ],
        'order': [{'column': 6, 'dir': 'desc'}, {'column': 5, 'dir': 'desc'}],
        'start': start, 'length': length,
        'search': {'value': '', 'regex': False},
        'additionalInfo': additional_info,
    })
    r = s.post(
        TAR_SEARCH, data=payload,
        headers={'Content-Type': 'application/json',
                 'X-Requested-With': 'XMLHttpRequest',
                 'Accept': 'application/json, text/javascript, */*',
                 'Referer': TAR_BASE},
        timeout=30, verify=False,
    )
    return r.json()

def scrape_tar():
    results = []
    print(f"\n{'='*60}\n[TAR Catanzaro] {TAR_BASE}")
    s = make_session({'Referer': 'https://www.giustizia-amministrativa.it/'})

    try:
        form_date = _tar_get_form_date(s)
        print(f"[TAR] form_date: {form_date}")
        _tar_post_form(s, form_date)

        start, length = 0, 200
        total_seen = 0
        while True:
            data = _tar_ajax_search(s, start=start, length=length)
            records_total    = data.get('recordsTotal', 0)
            records_filtered = data.get('recordsFiltered', 0)
            records          = data.get('data', [])
            print(f"[TAR] AJAX start={start}: total={records_total}, filtered={records_filtered}, returned={len(records)}")

            if not records:
                break

            for rec in records:
                total_seen += 1
                parte        = rec.get('parte', '') or ''
                tipo         = rec.get('tipoProvvedimento', '') or ''
                data_pub     = rec.get('dataPubblicazione', '') or ''
                num          = rec.get('numProvvedimento', '') or ''
                nrg          = rec.get('nrgFascicolo', '') or ''
                nome_file    = rec.get('nomeFile', '') or ''
                esito        = rec.get('esito', '') or ''
                relatore     = rec.get('relatore', '') or ''

                # Build document URL if possible
                doc_link = TAR_BASE
                if nome_file:
                    doc_link = (
                        f"https://www.giustizia-amministrativa.it/documentazione/sentenze-e-prassi/"
                        f"-/id/{num}"
                    )

                # Search text: primarily "parte", plus full row
                full_text = ' | '.join([parte, tipo, esito, relatore, nrg])
                if matches_keywords(full_text):
                    oggetto_val = f"{parte} – {tipo} n.{num} del {data_pub} (NRG {nrg}) Esito: {esito}"
                    print(f"[TAR] MATCH: {oggetto_val[:120]}")
                    results.append({'data': data_pub, 'oggetto': oggetto_val.strip(),
                                    'link': doc_link})

            if start + length >= records_filtered:
                break
            start += length
            time.sleep(0.3)

        print(f"[TAR] Scanned {total_seen} records, {len(results)} matches")

    except Exception as e:
        import traceback
        print(f"[TAR] ERROR: {e}")
        traceback.print_exc()

    return results

# ── Excel writer ──────────────────────────────────────────────────────────────

SHEET_NAMES = [
    'ASP Cosenza', 'ASP Catanzaro', 'ASP Crotone',
    'ASP Reggio Calabria', 'ASP Vibo Valentia',
    'Regione Calabria', 'TAR Catanzaro',
]

HEADER_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
LINK_FONT   = Font(color='0563C1', underline='single')
WRAP        = Alignment(wrap_text=True, vertical='top')
ERROR_FONT  = Font(color='C00000', italic=True)

def write_excel(all_data):
    """
    all_data: list of (results_list, error_msg_or_None) per sheet
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_name, (results, error_msg) in zip(SHEET_NAMES, all_data):
        ws = wb.create_sheet(title=sheet_name)
        ws.freeze_panes = 'A2'

        # Header row
        headers = ['Data', 'Oggetto', 'Descrizione breve (150 car.)', 'Link']
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font      = HEADER_FONT
            c.fill      = HEADER_FILL
            c.alignment = WRAP
        ws.row_dimensions[1].height = 18

        if error_msg:
            # Show error in red
            c = ws.cell(row=2, column=1, value=f"⚠ {error_msg}")
            c.font = ERROR_FONT
        elif not results:
            ws.cell(row=2, column=1, value='Nessun risultato nel periodo indicato')
        else:
            for ri, item in enumerate(results, 2):
                data_val    = item.get('data', '')
                oggetto_val = item.get('oggetto', '')
                link_val    = item.get('link', '')
                descr_val   = oggetto_val[:150] if oggetto_val else ''

                ws.cell(row=ri, column=1, value=data_val).alignment    = WRAP
                ws.cell(row=ri, column=2, value=oggetto_val).alignment = WRAP
                ws.cell(row=ri, column=3, value=descr_val).alignment   = WRAP

                lc           = ws.cell(row=ri, column=4, value=link_val)
                lc.alignment = WRAP
                if link_val and link_val.startswith('http'):
                    lc.hyperlink = link_val
                    lc.font      = LINK_FONT

        # Column widths
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 80
        ws.column_dimensions['C'].width = 55
        ws.column_dimensions['D'].width = 60

    wb.save(OUTPUT_FILE)
    print(f"\n✓ File salvato: {OUTPUT_FILE}")

# ── Main ──────────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    all_data = []

    # 5 ASP portals
    for name, url in ASP_PORTALS:
        results, err = scrape_asp(name, url)
        all_data.append((results, err))

    # Regione Calabria
    rc_results = scrape_regione_calabria()
    all_data.append((rc_results, None))

    # TAR Catanzaro
    tar_results = scrape_tar()
    all_data.append((tar_results, None))

    # Excel
    write_excel(all_data)

    # Summary
    print(f"\n{'='*60}")
    print("RIEPILOGO FINALE")
    for name, (res, err) in zip(SHEET_NAMES, all_data):
        if err:
            print(f"  {name:<25} ✗ {err}")
        else:
            print(f"  {name:<25} {len(res):>4} atti trovati")
    print(f"\nFile: {OUTPUT_FILE}")
